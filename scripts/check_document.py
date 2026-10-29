#!/usr/bin/env python3
"""Check a delivered HTML, XLSX, DOCX, PPTX, or Markdown file against references/file-design.html.

The checks follow the problem families in references/visual-design.html (删除测试 and 问题的同类判断), not a fixed
word list. The patterns still catch only common wordings, so a clean run is necessary, not sufficient:
the deletion test must be applied to every visible element by hand as well.
"""
from __future__ import annotations

import argparse
import re
import sys
from html.parser import HTMLParser
from pathlib import Path

# (pattern, family, short_only). short_only patterns fire only in short blocks (<= 80 characters), where
# metadata lines live, so the same words quoted inside long body text are not reported.
FAMILY_PATTERNS = [
    (r"\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}(?::\d{2})?(?:\.\d+)?Z?", "timestamp", True),
    (r"(?:生成|更新|修改|编制|导出|发布)(?:于|时间|日期)|最后(?:更新|修改)|(?:本(?:页面|页|文档|文|报告|表格|表|方案|清单)|此(?:页|文|表))[^。，,；;！？!?\n]{0,6}由[^。，,；;！？!?\n]{1,12}(?:生成|整理|撰写|编制|制作)|由\s*(?:AI|人工智能|Claude|ChatGPT|GPT|大模型|模型|智能体|AI\s?助手|Codex|Luz)[^。，,；;！？!?\n]{0,8}(?:生成|整理|撰写|编制|制作)|^\s*由[^。，,；;！？!?\n]{1,12}(?:生成|整理|撰写|编制|制作)\s*$|"
     r"版本[号:：]|编制人|制表人|审核人|作者[:：]|阅读(?:时长|时间)|字数[:：]|智能生成|人工智能生成|提示词|"
     r"Generated|Last updated|Created by|Powered by|Author:", "production metadata", False),
    (r"数据来源|信息来源|来源[:：]|去重来源|来源(?:统计|数量|索引|目录)|\d+\s*个(?:来源|官方页面|网页|链接)|官方页面|"
     r"采集(?:日期|时间|状态|范围)|证据(?:部分|状态|等级)|部分获取|(?:资料|信息|来源|数据|内容|事实)(?:已核验|核验通过|校验通过|已校验)|哈希|SHA-?256|"
     r"以官方(?:信息|公告)?为准|信息可能(?:不完整|有误|存在)|免责声明", "source metadata", False),
    (r"(?<![A-Za-z])raw/|单页\s*HTML|单页总结|本文件|[A-Za-z]:\\\S+|输出目录|源文件目录", "file metadata", False),
    (r"Design reference|设计参考|Luz Design|Admin SPEC|参考.{0,8}(?:规范|主题)(?:设计|制作)|"
     r"基于.{1,10}(?:主题|模板)(?:生成|制作)?|遵循.{1,12}规范", "design declaration", True),
    (r"本(?:报告|文档|文章|文|页面|页|节|章|表|图|部分|幻灯片|演示)(?:将|主要|旨在|用于|展示|介绍|列出|汇总|总结|说明|包含|分为|梳理)|"
     r"以下(?:是|为|内容|表格|列出|几点)|如下(?:表|图)?所示|如上所述|下面我们|我们来看|综上所述|总结一下|总而言之",
     "self-reference", False),
    (r"请注意|点击(?:展开|查看|这里|下方|此处)|向下滚动|阅读指南|使用方法[:：]|建议先(?:阅读|看)|详见下[页文]|"
     r"返回顶部|回到顶部", "reading guidance", False),
    (r"仅供参考|仅(?:作|为)示意|以实际为准|如需.{0,12}请|下一步(?:建议|计划|行动)|后续(?:可以|建议|可考虑)|可根据实际情况|建议(?:您|你|后续|进一步)|"
     r"欢迎(?:反馈|联系|交流|指正)|如有(?:疑问|问题|需要)", "advisory copy", False),
    (r"根据您的(?:要求|需求)|已为您|为您(?:整理|生成|准备|撰写)|希望(?:对您|能)?有(?:所)?帮助|以下为您|^好的[，,。]",
     "chat residue", False),
    (r"\bTODO\b|\bTBD\b|FIXME|lorem ipsum|\bXXX\b", "placeholder", False),
    (r"待补充|待完善|待填写|示例数据|占位文字|（略）|\(略\)", "placeholder", True),
    (r"全网最|最全|史上最|高质量(?:报告|内容|文档)|权威(?:指南|报告)|深度(?:好文|解析)|震撼", "self-evaluation", False),
]
FAMILY_PATTERNS = [(re.compile(p, re.I | re.M), family, short) for p, family, short in FAMILY_PATTERNS]
FAMILIES = sorted({family for _, family, _ in FAMILY_PATTERNS} | {"note", "emoji"})
NOTE_PREFIX = re.compile(r"^\s*(?:说明|注|注意|提示|备注|补充|温馨提示|小贴士|注意事项|建议|Tips?|PS|Note)\s*[:：]", re.I)
EMOJI = re.compile(r"[\U0001F300-\U0001FAFF☀-➿⭐⭕■-◿]")
# Unrequested accessory sections (file-design.html, 同类问题 for documents and slides).
ACCESSORY_HEADING = re.compile(
    r"^(?:\S{0,4}[、.．]?\s*)?(?:摘要|前言|序言|编写(?:说明|目的)|阅读(?:指南|说明)|使用说明|修订记录|版本(?:信息|记录|历史)|"
    r"更新(?:记录|日志)|免责声明|致谢|术语表|关于本(?:文档|报告|文)|本章小结|小结|参考来源|数据来源|附录说明)$")
FILLER_SLIDE = re.compile(r"^(?:目录|议程|大纲|内容提要|谢谢|感谢|谢谢观看|感谢聆听|问答|答疑|结束|结语|关于我们|附录|Q\s*&\s*A|Thanks?|Agenda)[!！。]*$", re.I)
GRAPHIC_SELECTOR = re.compile(r"svg|\.icon|frame|illustration|\.empty|chart|star|\.kk|\.logo|face|ring|scan", re.I)
DARK_SMALL_SELECTOR = re.compile(r"\.(?:dot|icon|legend|bar|badge|pill|status|progress|chip|tag|logo|corner|coins?|kk|steps|mark|avatar|check|box|ring|scan|cbox|heat|rb)\b|-(?:mark|bar)\b|\.copy i|::?(?:before|after|marker)|svg|button|input", re.I)
# Filled controls in their active state: primary and danger buttons, checked switches and boxes, selected options.
ACTIVE_CONTROL_SELECTOR = re.compile(r"primary|danger|aria-checked=.?true|aria-pressed=.?true|\.on\b|ck-mark", re.I)
# The one brand card per screen (visual-design.html 配色, 品牌主卡): a solid primary-step card with white text.
BRAND_CARD = re.compile(r"\.brand-card(?![\w-])")
BRAND_CARD_CLASS = "brand-card"
BRAND_CARD_MIN_CONTRAST = 4.5
# One line icon set (visual-design.html 图形：一套线性图标): 2 on the 24 grid, 1.75 allowed for 16px icons.
ICON_SELECTOR = re.compile(r"\.icon\b|\.icon-|svg\.fi\b|\bfi\b|svg\.i\b|\.ico\b")
ICON_STROKES = {1.75, 2.0}
# Three weights (visual-design.html 字体与字号): 400 body, 600 headings and emphasis, 700 only at 22px and larger.
PLAIN_WEIGHTS = {"400", "600", "normal", "inherit", "initial", "unset"}
BOLD_WEIGHTS = {"700", "bold"}
BOLD_MIN_PX = 22.0
# Elevation (visual-design.html 层级与立体感): cards take levels 1–2 (far blur up to 24), overlays up to level 4 (48).
CARD_MAX_BLUR = 24.0
OVERLAY_MAX_BLUR = 48.0
SHADOW_MAX_ALPHA = 0.2
ALLOWED: set[str] = set()


def FLOOR_PX() -> float:
    """Reading pages keep 14px; interface prototypes (--allow interface) use the 12px annotation step."""
    return 12.0 if "interface" in ALLOWED else 14.0


# One shared scale for the function list, mini-program and admin of a product deliverable (visual-design.html, 字阶).
INTERFACE_SCALE = {12, 14, 16, 18, 22, 28, 32}
# Websites add the first-screen title steps (website-design.html, 六、视觉基础): 40 on desktop, 48 for a brand home page.
WEBSITE_SCALE = INTERFACE_SCALE | {40, 48}


def luminance(rgb: tuple[int, int, int]) -> float:
    def channel(value: int) -> float:
        v = value / 255
        return v / 12.92 if v <= 0.04045 else ((v + 0.055) / 1.055) ** 2.4
    r, g, b = (channel(v) for v in rgb)
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def hex_rgb(value: str) -> tuple[int, int, int]:
    value = value[-6:]
    return tuple(int(value[i:i + 2], 16) for i in (0, 2, 4))


def parse_colors_alpha(value: str) -> list[tuple[tuple[int, int, int], float]]:
    """Every colour in a value with its alpha (hex and rgb are opaque)."""
    found = []
    for hex_value in re.findall(r"#([0-9a-fA-F]{6}|[0-9a-fA-F]{3})\b", value):
        if len(hex_value) == 3:
            hex_value = "".join(c * 2 for c in hex_value)
        found.append((hex_rgb(hex_value), 1.0))
    for match in re.finditer(r"rgba?\(\s*(\d+)[ ,]+(\d+)[ ,]+(\d+)(?:[ ,/]+([\d.]+%?))?", value):
        alpha = match.group(4)
        a = 1.0 if not alpha else float(alpha.rstrip("%")) / (100 if alpha.endswith("%") else 1)
        found.append((tuple(int(match.group(i)) for i in (1, 2, 3)), a))
    return found


def shadow_layers(value: str) -> list[dict]:
    """Each layer of a box-shadow or text-shadow: inset, x, y, blur, spread, colour and alpha (no colour = opaque black)."""
    layers = []
    if value.strip().lower() in {"none", "0", "inherit", "initial", "unset", ""}:
        return layers
    for part in re.split(r",(?![^(]*\))", value):
        if not part.strip():
            continue
        lengths = re.sub(r"(?:rgba?|hsla?|var)\([^)]*\)|#[0-9a-f]{3,8}\b|\binset\b", " ", part, flags=re.I)
        nums = [float(n) for n in re.findall(r"(-?\d*\.?\d+)(?:px)?\b", lengths)] + [0.0, 0.0, 0.0, 0.0]
        colours = parse_colors_alpha(part)
        colour, alpha = colours[0] if colours else ((0, 0, 0), 1.0)
        layers.append({"inset": bool(re.search(r"\binset\b", part, re.I)), "x": nums[0], "y": nums[1], "blur": nums[2],
                       "spread": nums[3], "colour": colour, "alpha": alpha, "text": part.strip()[:34]})
    return layers


def blend(color: tuple[int, int, int], alpha: float, backdrop: tuple[int, int, int]) -> tuple[int, int, int]:
    return tuple(round(c * alpha + b * (1 - alpha)) for c, b in zip(color, backdrop))


def contrast_ratio(a: tuple[int, int, int], b: tuple[int, int, int]) -> float:
    lighter, darker = sorted((luminance(a), luminance(b)), reverse=True)
    return (lighter + 0.05) / (darker + 0.05)


# Ant Design official 10-step columns (ant.design/docs/spec/colors-cn, step 1 = lightest). A theme that takes one step of a
# column may use the column's other steps in icons and illustrations.
ANT_COLUMNS = {
    "red": "#FFF1F0 #FFCCC7 #FFA39E #FF7875 #FF4D4F #F5222D #CF1322 #A8071A #820014 #5C0011",
    "volcano": "#FFF2E8 #FFD8BF #FFBB96 #FF9C6E #FF7A45 #FA541C #D4380D #AD2102 #871400 #610B00",
    "orange": "#FFF7E6 #FFE7BA #FFD591 #FFC069 #FFA940 #FA8C16 #D46B08 #AD4E00 #873800 #612500",
    "lime": "#FCFFE6 #F4FFB8 #EAFF8F #D3F261 #BAE637 #A0D911 #7CB305 #5B8C00 #3F6600 #254000",
    "gold": "#FFFBE6 #FFF1B8 #FFE58F #FFD666 #FFC53D #FAAD14 #D48806 #AD6800 #874D00 #613400",
    "yellow": "#FEFFE6 #FFFFB8 #FFFB8F #FFF566 #FFEC3D #FADB14 #D4B106 #AD8B00 #876800 #614700",
    "green": "#F6FFED #D9F7BE #B7EB8F #95DE64 #73D13D #52C41A #389E0D #237804 #135200 #092B00",
    "cyan": "#E6FFFB #B5F5EC #87E8DE #5CDBD3 #36CFC9 #13C2C2 #08979C #006D75 #00474F #002329",
    "blue": "#E6F4FF #BAE0FF #91CAFF #69B1FF #4096FF #1677FF #0958D9 #003EB3 #002C8C #001D66",
    "geekblue": "#F0F5FF #D6E4FF #ADC6FF #85A5FF #597EF7 #2F54EB #1D39C4 #10239E #061178 #030852",
    "purple": "#F9F0FF #EFDBFF #D3ADF7 #B37FEB #9254DE #722ED1 #531DAB #391085 #22075E #120338",
    "magenta": "#FFF0F6 #FFD6E7 #FFADD2 #FF85C0 #F759AB #EB2F96 #C41D7F #9E1068 #780650 #520339",
}


def theme_palette() -> set[tuple[int, int, int]]:
    """Every colour defined in references/themes.json (shared tokens, theme steps, scales and module accents), plus the other
    steps of each official Ant Design column a theme already uses."""
    import json
    path = Path(__file__).resolve().parents[1] / "references" / "themes.json"
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return set()
    text = json.dumps({"shared": data.get("shared", {}),
                       "themes": {k: {x: y for x, y in v.items() if x != "dark"} for k, v in data.get("themes", {}).items()}})
    colours = {hex_rgb(h) for h in re.findall(r"#([0-9a-fA-F]{6})\b", text)}
    for column in ANT_COLUMNS.values():
        steps = {hex_rgb(h) for h in re.findall(r"#([0-9A-F]{6})", column)}
        if steps & colours:
            colours |= steps
    return colours


def light_neutral(color: tuple[int, int, int]) -> bool:
    return max(color) - min(color) <= 8 and luminance(color) >= 0.6


def template_styles() -> list[tuple[str, str]]:
    """(name, normalised <style>) for the HTML templates in references/templates.txt."""
    path = Path(__file__).resolve().parents[1] / "references" / "templates.txt"
    try:
        text = path.read_text(encoding="utf-8")
    except OSError:
        return []
    found = []
    for name, body in re.findall(r"=== template: ([\w.-]+\.html) ===\n(.*?)\n=== end ===", text, re.S):
        style = re.search(r"<style[^>]*>(.*?)</style>", body, re.S | re.I)
        if style:
            found.append((name, re.sub(r"\s+", "", style.group(1))))
    return found


def parse_colors(value: str) -> list[tuple[int, int, int]]:
    colors = []
    for hex_value in re.findall(r"#([0-9a-fA-F]{6}|[0-9a-fA-F]{3})\b", value):
        if len(hex_value) == 3:
            hex_value = "".join(c * 2 for c in hex_value)
        colors.append(hex_rgb(hex_value))
    for match in re.finditer(r"rgba?\(\s*(\d+)[ ,]+(\d+)[ ,]+(\d+)(?:[ ,/]+([\d.]+%?))?", value):
        alpha = match.group(4)
        if alpha:
            a = float(alpha.rstrip("%")) / (100 if alpha.endswith("%") else 1)
            if a < 0.5:
                continue
        colors.append(tuple(int(match.group(i)) for i in (1, 2, 3)))
    return colors


def check_text(text: str, where: str, errors: list[str]) -> None:
    for pattern, family, short_only in FAMILY_PATTERNS:
        if family in ALLOWED or (short_only and len(text) > 80):
            continue
        match = pattern.search(text)
        if match:
            errors.append(f"{where}: banned {family}: {match.group(0)!r}")
    if "note" not in ALLOWED and NOTE_PREFIX.search(text):
        errors.append(f"{where}: note or advice prefix: {text.strip()[:40]!r}")
    if "emoji" not in ALLOWED and EMOJI.search(text):
        errors.append(f"{where}: emoji or decorative symbol: {EMOJI.search(text).group(0)!r}")


def check_heading(text: str, where: str, errors: list[str]) -> None:
    if "accessory" not in ALLOWED and ACCESSORY_HEADING.match(text.strip()):
        errors.append(f"{where}: unrequested accessory section {text.strip()!r} (pass --allow accessory if the user asked for it)")


# ---------------------------------------------------------------- HTML

class Outline(HTMLParser):
    """Collect visible text blocks in document order with their nearest block tag and ancestors."""

    HIDDEN = {"script", "style", "template", "head", "title", "noscript"}
    VOID = {"br", "img", "meta", "link", "input", "hr", "source", "wbr", "col", "path", "circle", "rect",
            "line", "polyline", "polygon", "ellipse", "use"}
    BLOCKS = {"h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "td", "th", "div", "span", "a", "figcaption",
              "caption", "summary", "dt", "dd", "header", "nav", "aside", "footer", "time", "strong", "em",
              "small", "label", "button", "section", "article", "main", "blockquote", "pre", "code"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.stack: list[tuple[str, bool, bool]] = []
        self.blocks: list[tuple[str, str, tuple[str, ...], bool]] = []
        self.styles: list[str] = []
        self.in_style = False

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag in self.VOID:
            if attrs.get("style"):
                self.styles.append("x{" + attrs["style"] + "}")
            return
        hidden = (
            tag in self.HIDDEN
            or "hidden" in attrs
            or (attrs.get("aria-hidden") or "").lower() == "true"
            or bool(re.search(r"display\s*:\s*none", attrs.get("style") or "", re.I))
        )
        # Verbatim third-party text (captured pages, comments, citations) is marked with data-quoted or <blockquote>.
        quoted = "data-quoted" in attrs or tag in {"blockquote", "q"} or bool(self.stack and self.stack[-1][2])
        self.stack.append((tag, hidden or bool(self.stack and self.stack[-1][1]), quoted))
        if tag == "style":
            self.in_style = True
        if attrs.get("style"):
            self.styles.append("x{" + attrs["style"] + "}")

    def handle_endtag(self, tag):
        if tag == "style":
            self.in_style = False
        for index in range(len(self.stack) - 1, -1, -1):
            if self.stack[index][0] == tag:
                del self.stack[index:]
                break

    def handle_data(self, data):
        if self.in_style:
            self.styles.append(data)
            return
        if not data.strip() or (self.stack and self.stack[-1][1]):
            return
        tag = next((t for t, _, _ in reversed(self.stack) if t in self.BLOCKS), "text")
        quoted = bool(self.stack and self.stack[-1][2])
        self.blocks.append((tag, " ".join(data.split()), tuple(t for t, _, _ in self.stack), quoted))


def is_scrim(value: str) -> bool:
    """A photo scrim: every stop is black or white with alpha, or transparent."""
    stops = re.findall(r"#[0-9a-f]{3,8}\b|rgba?\([^)]*\)|\btransparent\b|var\([^)]*\)", value, re.I)
    return bool(stops) and all(re.fullmatch(r"transparent|rgba\(\s*(?:0\s*,\s*0\s*,\s*0|255\s*,\s*255\s*,\s*255)\s*,[^)]*\)", x.strip(), re.I) for x in stops)


def is_hard_stop(value: str) -> bool:
    """A hard-stop fill such as a slider track: two consecutive stops share one position."""
    positions = re.findall(r"\s(\d+(?:\.\d+)?%|var\([^)]*\))\s*(?=[,)])", value)
    return any(a == b for a, b in zip(positions, positions[1:]))


# ---------------------------------------------------------------- Tab icon (every HTML mode)
# Every delivered HTML carries its own browser-tab icon: <link rel="icon" href="data:image/svg+xml,...">, a flat line
# pictogram of the page's subject in the theme's primary colour on a rounded tint square (visual-design.html 网页图标与品牌标识;
# file-design.html, HTML 必备项). Missing, empty or an unfilled template placeholder is an error.
ICON_DRAWING = re.compile(r"<(?:path|rect|circle|line|polyline|polygon|ellipse)\b", re.I)


class _IconLinks(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.links: list[str] = []

    def handle_starttag(self, tag, attrs):
        if tag == "link":
            values = {k.lower(): (v or "") for k, v in attrs}
            if "icon" in values.get("rel", "").lower():
                self.links.append(values.get("href", "").strip())

    handle_startendtag = handle_starttag


def icon_svg(href: str) -> str | None:
    """The decoded SVG of a data:image/svg+xml icon, or None when the href is not an inline SVG."""
    from base64 import b64decode
    from urllib.parse import unquote
    if not href.lower().startswith("data:") or "," not in href:
        return None
    meta, payload = href[5:].split(",", 1)
    if "svg" not in meta.lower():
        return None
    try:
        return b64decode(payload).decode("utf-8", "replace") if ";base64" in meta.lower() else unquote(payload)
    except ValueError:
        return ""


def tab_icon_findings(source: str) -> tuple[list[str], list[str]]:
    """(errors, warnings) for the page's browser-tab icon."""
    parser = _IconLinks()
    parser.feed(source)
    hint = ('add <link rel="icon" href="data:image/svg+xml,..."> — the theme pictogram in the primary colour on a rounded '
            "tint square (visual-design.html 网页图标与品牌标识)")
    if not parser.links:
        return [f"no browser-tab icon: {hint}"], []
    usable = []
    for href in parser.links:
        svg = icon_svg(href)
        empty = (not href or href in {"#", "about:blank"} or "{{" in href or re.fullmatch(r"data:[^,]*,\s*", href, re.I)
                 or (svg is not None and not ICON_DRAWING.search(svg)))
        if not empty:
            usable.append((href, svg))
    if not usable:
        return [f"browser-tab icon is empty or an unfilled placeholder: {hint}"], []
    warnings = []
    href, svg = usable[0]
    if svg is None and not href.lower().startswith("data:"):
        warnings.append(f"browser-tab icon {href[:40]!r} is an external file: inline it as data:image/svg+xml so the page carries its own icon")
    if svg is not None and re.search(r"<text\b", svg, re.I):
        warnings.append("browser-tab icon contains text: draw a pictogram, never letters or initials")
    return [], warnings


def check_html(path: Path, errors: list[str], warnings: list[str]) -> None:
    source = path.read_text(encoding="utf-8")
    icon_errors, icon_warnings = tab_icon_findings(source)
    errors.extend(icon_errors)
    warnings.extend(icon_warnings)
    outline = Outline()
    outline.feed(source)
    blocks = outline.blocks
    if not blocks:
        errors.append("no visible text")
        return
    if blocks[0][0] != "h1":
        errors.append(f"first visible text is not the h1 title: <{blocks[0][0]}> {blocks[0][1][:40]!r}")
    h1_count = len(re.findall(r"<h1[\s>]", source, re.I))
    if h1_count != 1:
        errors.append(f"expected exactly one h1, found {h1_count}")
    tags = [block[0] for block in blocks]
    if "h1" in tags:
        after = tags.index("h1") + 1
        while after < len(tags) and tags[after] == "h1":
            after += 1
        if after < len(tags) and "h2" in tags and tags[after] != "h2":
            errors.append(f"text between the h1 and the first h2: {blocks[after][1][:40]!r}")

    shell_reported: set[str] = set()
    for tag, text, ancestors, quoted in blocks:
        where = f"<{tag}>"
        if not quoted:
            check_text(text, where, errors)
            if tag in {"h2", "h3", "h4"}:
                check_heading(text, where, errors)
        in_content = any(a in {"article", "section"} for a in ancestors)
        shell = next((a for a in ancestors if a in {"nav", "aside"} or (a in {"header", "footer"} and not in_content)), None)
        # An interface prototype is an application: its shell (nav, side bar, top bar) belongs to it.
        if shell and not ALLOWED & {"shell", "interface"} and shell not in shell_reported:
            shell_reported.add(shell)
            errors.append(f"application shell <{shell}> in a document: {text[:30]!r}")
    for index, (tag, text, _, _) in enumerate(blocks[:-2]):
        nxt_tag, nxt_text, _, _ = blocks[index + 1]
        if tag in {"h2", "h3"} and nxt_tag == "p" and len(nxt_text) <= 40 and blocks[index + 2][0] in {"table", "th", "td", "li", "h3", "h4", "div"}:
            warnings.append(f"possible explanatory line under heading {text[:20]!r}: {nxt_text!r}")

    css = "\n".join(outline.styles)
    variables = dict(re.findall(r"(--[\w-]+)\s*:\s*([^;{}]+)", css))

    def resolve(value: str, depth: int = 0) -> str:
        if depth > 8:
            return value
        return re.sub(r"var\((--[\w-]+)(?:\s*,\s*([^)]+))?\)",
                      lambda m: resolve(variables.get(m.group(1), m.group(2) or ""), depth + 1), value)

    if re.search(r"color-scheme\s*:\s*dark|prefers-color-scheme\s*:\s*dark", css, re.I) and "dark" not in ALLOWED:
        errors.append("dark color scheme declared")
    families: set[str] = set()
    icon_strokes: set[str] = set()
    bold_rules: list[tuple[str, str]] = []
    # Brand cards (visual-design.html 配色, 品牌主卡): their fill is the backdrop for text in descendant rules.
    brand_fills: list[tuple[int, int, int]] = []
    for rule in re.finditer(r"([^{}]*)\{([^{}]*)\}", css):
        sel, bod = rule.group(1).strip(), rule.group(2)
        if any(BRAND_CARD.search(part.split()[-1]) for part in sel.split(",") if part.split()):
            brand_fills += [c for m in re.finditer(r"(?:^|;)\s*background(?:-color)?\s*:\s*([^;]+)", bod, re.I)
                            for c in parse_colors(resolve(m.group(1))) if luminance(c) < 0.6]
    for rule in re.finditer(r"([^{}]*)\{([^{}]*)\}", css):
        selector, body = rule.group(1).strip(), rule.group(2)
        parts = [part.split() for part in selector.split(",") if part.split()]
        brand_self = bool(parts) and all(BRAND_CARD.search(part[-1]) for part in parts)
        brand_inside = bool(parts) and not brand_self and all(any(BRAND_CARD.search(x) for x in part[:-1]) for part in parts)
        for decl in re.finditer(r"(?:^|;)\s*font-size\s*:\s*([^;]+)", body, re.I):
            value = resolve(decl.group(1))
            for number, unit in re.findall(r"(\d+(?:\.\d+)?)\s*(px|rem|em|pt)\b", value, re.I):
                if float(number) * {"px": 1, "rem": 16, "em": 16, "pt": 4 / 3}[unit.lower()] < FLOOR_PX():
                    errors.append(f"font size below {FLOOR_PX():g}px in {selector[:40]!r}: {decl.group(1).strip()}")
                elif "interface" in ALLOWED and unit.lower() == "px" and "calc" not in value and float(number) not in (WEBSITE_SCALE if "website" in ALLOWED else INTERFACE_SCALE):
                    errors.append(f"font size off the product type scale 12/14/16/18/22/28/32 in {selector[:40]!r}: {decl.group(1).strip()}")
        if "interface" in ALLOWED and re.search(r"drawer|side-?panel|slide-?over", selector, re.I) and not re.search(r"menu|nav", selector, re.I):
            # Details and forms float in the centre; only navigation may slide in from the side (Luz-System ui-standards.html, 浮层与反馈).
            errors.append(f"side drawer for content in {selector[:40]!r}: use a centred floating panel")
        if "interface" in ALLOWED:
            # Product deliverables use three weights and no tracking (visual-design.html, 字体与字号): 400 body,
            # 600 headings and emphasis, 700 only for page titles and large numbers at 22px and larger.
            for decl in re.finditer(r"(?:^|;)\s*font-weight\s*:\s*([^;]+)", body, re.I):
                weight = decl.group(1).replace("!important", "").strip().lower()
                if weight in BOLD_WEIGHTS:
                    if selector != "x":
                        bold_rules.append((selector, body))
                elif weight not in PLAIN_WEIGHTS:
                    errors.append(f"font weight outside 400/600/700 in {selector[:40]!r}: {decl.group(1).strip()}")
            for decl in re.finditer(r"(?:^|;)\s*letter-spacing\s*:\s*([^;]+)", body, re.I):
                if decl.group(1).strip().lower() not in {"0", "normal", "0px", "inherit"}:
                    errors.append(f"letter spacing in {selector[:40]!r}: {decl.group(1).strip()}")
        else:
            # Reading pages use the same three weights (file-design.html 网页文件): 700 only for the 24 and 32px titles.
            own_size = font_size_value(declarations(body, resolve))
            own_px = length_px(own_size, 16.0) if own_size else None
            for decl in re.finditer(r"(?:^|;)\s*font-weight\s*:\s*([^;]+)", body, re.I):
                weight = decl.group(1).replace("!important", "").strip().lower()
                if weight in BOLD_WEIGHTS and own_px is not None and own_px < BOLD_MIN_PX:
                    warnings.append(f"font weight 700 at {own_px:g}px in {selector[:40]!r}: below 22px use 600")
                elif weight not in BOLD_WEIGHTS | PLAIN_WEIGHTS:
                    warnings.append(f"font weight outside 400/600/700 in {selector[:40]!r}: {decl.group(1).strip()}")
        if re.search(r"(?:^|;)\s*font-style\s*:\s*(?:italic|oblique)", body, re.I):
            errors.append(f"italic text in {selector[:40]!r}: Chinese has no true italic")
        for decl in re.finditer(r"(?:^|;)\s*font\s*:\s*([^;]+)", body, re.I):
            size = re.search(r"(\d+(?:\.\d+)?)px", resolve(decl.group(1)))
            if size and float(size.group(1)) < FLOOR_PX():
                errors.append(f"font size below {FLOOR_PX():g}px in {selector[:40]!r}: {decl.group(1).strip()[:30]}")
        for decl in re.finditer(r"(?:^|;)\s*(?:font-family|font)\s*:\s*([^;]+)", body, re.I):
            family = resolve(decl.group(1))
            # In the font shorthand the family list follows the size (and optional /line-height).
            family_list = re.sub(r"^.*?\d+(?:\.\d+)?(?:px|pt|em|rem|%)(?:\s*/\s*[\d.]+\w*)?\s+", "", family.strip()) \
                if decl.group(0).lower().lstrip("; ").startswith("font:") else family
            token = family_list.split(",")[0].strip().strip("\"'").strip()
            first = "" if not token or token.lower() in {"inherit", "initial", "unset"} or token.startswith("var(") \
                or re.fullmatch(r"[\d.]+\w*", token) else token
            if first and not re.search(r"mono|consol|courier|code", family, re.I) and not re.search(r"\b(?:pre|code|kbd|samp)\b", selector):
                families.add(first.lower())
        for decl in re.finditer(r"(?:^|;)\s*background(?:-color|-image)?\s*:\s*([^;]+)", body, re.I):
            value = resolve(decl.group(1))
            small = DARK_SMALL_SELECTOR.search(selector)
            # Flat style everywhere (visual-design.html, 扁平线性风格): no decorative gradients. Interfaces may keep a
            # black-alpha scrim over a photo, a hard-stop fill (sliders, progress) and skeleton shimmer.
            if re.search(r"gradient", value, re.I):
                if "interface" not in ALLOWED:
                    if not small:
                        errors.append(f"gradient background in {selector[:40]!r}")
                elif not (is_scrim(value) or is_hard_stop(value) or re.search(r"\bsk\b|\.ask\b|skeleton|shimmer", selector, re.I)):
                    errors.append(f"decorative gradient in {selector[:40]!r}: use one flat colour (扁平线性风格)")
            for color in parse_colors(value):
                lum = luminance(color)
                if lum < 0.6 and brand_self and "interface" in ALLOWED:
                    # One brand card per screen may take a solid primary step: a hue, not a dark neutral, with white
                    # (or the declared) text at 4.5:1 (visual-design.html 配色, 品牌主卡; WCAG 1.4.3).
                    texts = [c for m in re.finditer(r"(?:^|;)\s*color\s*:\s*([^;]+)", body, re.I) for c in parse_colors(resolve(m.group(1)))]
                    ratio = contrast_ratio(texts[0] if texts else (255, 255, 255), color)
                    if max(color) - min(color) < 60:
                        errors.append(f"brand card {selector[:40]!r} is a dark neutral {decl.group(1).strip()[:30]}: use the primary brand or pressed step")
                    elif ratio < BRAND_CARD_MIN_CONTRAST:
                        errors.append(f"brand card {selector[:40]!r}: text on {decl.group(1).strip()[:30]} is {ratio:.1f}:1, "
                                      f"needs {BRAND_CARD_MIN_CONTRAST:g}:1 (use the step where white text reaches 4.5:1)")
                elif lum < 0.6 and not small and "dark" not in ALLOWED:
                    errors.append(f"dark or saturated fill in {selector[:40]!r}: {decl.group(1).strip()[:30]} (luminance {lum:.2f})")
                elif ("interface" in ALLOWED and 0.3 < lum < 0.7 and ACTIVE_CONTROL_SELECTOR.search(selector)
                      and not re.search(r"::?(?:before|after)|-thumb|:disabled|:hover", selector)
                      and (1.05 / (lum + 0.05)) < 3.0):
                    # A coloured control fill that is neither a pale container nor a solid step fails WCAG 1.4.11 (3:1)
                    # and reads as washed-out (visual-design.html, 取色依据).
                    errors.append(f"washed-out control fill in {selector[:40]!r}: {decl.group(1).strip()[:30]} ({1.05 / (lum + 0.05):.1f}:1 against white, needs 3:1 or a pale container)")
            if ("interface" in ALLOWED and re.search(r"gradient", value, re.I) and re.search(r"primary|danger", selector, re.I)
                    and not re.search(r"::?(?:before|after)|-thumb|prob", selector)):
                # Design systems fill buttons with one flat palette step (Ant Design, TDesign, Material 3).
                errors.append(f"gradient on control in {selector[:40]!r}: use one flat palette step")
        own_bg = [c for m in re.finditer(r"(?:^|;)\s*background(?:-color)?\s*:\s*([^;]+)", body, re.I) for c in parse_colors(resolve(m.group(1)))]
        # Text inside a brand card sits on the card's solid fill, not on white.
        inside_card = brand_inside and brand_fills and "interface" in ALLOWED
        backdrop = own_bg[0] if own_bg else (brand_fills[0] if inside_card else (255, 255, 255))
        # Text on a transparent layer sits over imagery or a parent surface we cannot see here; skip it.
        transparent_layer = bool(re.search(r"background(?:-color)?\s*:\s*(?:transparent|none)\b", body, re.I))
        # Text needs 4.5:1; graphics (icons, frames, illustrations) and disabled controls need 3:1.
        minimum = 3.0 if GRAPHIC_SELECTOR.search(selector) or ":disabled" in selector else 4.5
        for decl in (() if transparent_layer else re.finditer(r"(?:^|;)\s*color\s*:\s*([^;]+)", body, re.I)):
            for color, alpha in parse_colors_alpha(resolve(decl.group(1))):
                if alpha < 0.5 and not inside_card:
                    continue
                shown = blend(color, alpha, backdrop) if alpha < 1 else color
                contrast = contrast_ratio(shown, backdrop)
                if contrast < minimum:
                    errors.append(f"color too light in {selector[:40]!r}: {decl.group(1).strip()} ({contrast:.1f}:1, needs {minimum:g}:1)")
        if "interface" not in ALLOWED:
            # Reading pages stay on level 0 with hairlines (visual-design.html 层级与立体感; file-design.html 网页文件).
            for decl in re.finditer(r"(?:^|;)\s*(?:box-shadow|text-shadow)\s*:\s*([^;]+)", body, re.I):
                blurred = [layer for layer in shadow_layers(resolve(decl.group(1)))
                           if not layer["inset"] and layer["blur"] >= 4 and layer["alpha"] > 0]
                if blurred:
                    warnings.append(f"shadow in a reading page {selector[:40]!r}: {decl.group(1).strip()[:30]} "
                                    "(reading pages use level 0, hairlines and pale fills; no card shadows)")
        if "interface" in ALLOWED:
            # Elevation (visual-design.html 层级与立体感; Material 3 elevation, Fluent 2 shadows, Ant Design shadow tokens):
            # two-layer neutral (or dark theme-tinted) shadows by level; cards up to level 2, overlays up to level 4.
            # 1–2px edges (an inset top highlight, a button's pressed-step bottom edge) and focus rings are lines, not bevels.
            overlay = re.search(r"modal|dialog|pop|toast|tip|drop|sheet|layer|confirm|drawer|menu|float|picker|lightbox|panel|sticky|fixed",
                                selector, re.I)
            # Bars docked at the bottom mirror their level upward (Ant Design upward shadows for bottom navigation and toolbars).
            bottom_bar = re.search(r"bottom|tab-?bar|foot|action-?bar|toolbar|dock", selector, re.I)
            for decl in re.finditer(r"(?:^|;)\s*box-shadow\s*:\s*([^;]+)", body, re.I):
                for layer in shadow_layers(resolve(decl.group(1))):
                    inset, x, y, blur, spread = layer["inset"], layer["x"], layer["y"], layer["blur"], layer["spread"]
                    colour, alpha, text = layer["colour"], layer["alpha"], layer["text"]
                    bright = luminance(colour) > 0.5
                    # A primary or other mid-tone hue casts a coloured shadow; a dark theme tint (the 10th step) is allowed.
                    tinted = max(colour) - min(colour) > 40 and luminance(colour) > 0.05
                    if alpha == 0:
                        continue
                    if inset:
                        edge = x == 0 and abs(y) <= 2 and blur <= 2 and spread <= 0
                        ring = x == 0 and y == 0 and blur == 0 and 0 < spread <= 2
                        if not (edge or ring):
                            errors.append(f"bevel shadow in {selector[:40]!r}: {text} (only a 1–2px inset edge or ring; no raised or sunken bevels)")
                    elif blur <= 1 and abs(y) <= 2 and x == 0 and spread <= 0:
                        continue  # a 1–2px edge line, such as a button's pressed-step bottom edge
                    elif blur == 0 and x == 0 and y == 0:
                        continue  # a solid ring (focus ring, halo around a round control)
                    elif x == 0 and y == 0 and blur >= 8 and (bright or tinted or alpha > SHADOW_MAX_ALPHA):
                        errors.append(f"glow in {selector[:40]!r}: {text}")
                    elif (bright or tinted) and alpha >= 0.1 and blur >= 4:
                        errors.append(f"coloured shadow in {selector[:40]!r}: {text} (shadows are neutral or a dark theme tint)")
                    elif y < 0 and blur >= 2 and not bottom_bar:
                        warnings.append(f"upward shadow in {selector[:40]!r}: {text} (one light source above: shadows fall downward; "
                                        "only bars docked at the bottom mirror the level upward)")
                    elif blur > (OVERLAY_MAX_BLUR if overlay else CARD_MAX_BLUR):
                        warnings.append(f"shadow beyond the elevation levels in {selector[:40]!r}: {text} "
                                        f"(cards use levels 1–2, far blur up to {CARD_MAX_BLUR:g}; overlays up to {OVERLAY_MAX_BLUR:g})")
                    elif alpha > SHADOW_MAX_ALPHA and blur >= 4:
                        warnings.append(f"dark shadow in {selector[:40]!r}: {text} (layer alpha up to {SHADOW_MAX_ALPHA:g})")
            if re.search(r"(?:^|;)\s*(?:-webkit-)?backdrop-filter\s*:\s*[^;]*blur", body, re.I):
                errors.append(f"glass blur in {selector[:40]!r}: flat surfaces are opaque")
            if re.search(r"(?:^|;)\s*border-radius\s*:\s*[^;]*\d+%\s+\d+%[^;]*", body, re.I):
                errors.append(f"organic blob shape in {selector[:40]!r}")
            if ICON_SELECTOR.search(selector) and not re.search(r"chart|check|box|avatar|art\b", selector):
                for decl in re.finditer(r"(?:^|;)\s*stroke-width\s*:\s*([\d.]+)", body, re.I):
                    icon_strokes.add(decl.group(1))
    if len(families) > 1:
        errors.append(f"more than one font family: {sorted(families)}")
    # One typeface everywhere (visual-design.html 字体与字号): the canonical stack is
    # "PingFang SC", "Microsoft YaHei", "Noto Sans SC", system-ui, sans-serif; a page may lead with any of those three.
    allowed_first = {"pingfang sc", "microsoft yahei", "noto sans sc", "苹方", "微软雅黑", "思源黑体", "system-ui", "sans-serif", "inherit"}
    for family in sorted(families - allowed_first):
        errors.append(f"font family {family!r} is not the shared Chinese sans stack (PingFang SC / Microsoft YaHei / Noto Sans SC)")
    # Sizes outside <style>: inline style attributes, SVG text attributes and sizes written into scripts.
    style_text = re.sub(r"\s+", "", "\n".join(outline.styles))
    for number, unit in re.findall(r"font-size\s*[:=]\s*[\"']?\s*(\d+(?:\.\d+)?)(px)?(?![\d.]*(?:pt|em|rem|%))", source, re.I):
        if float(number) < FLOOR_PX() and f"font-size:{number}" not in style_text:
            errors.append(f"font size {number}{unit or ''} below {FLOOR_PX():g}px in an inline style, SVG attribute or script")
            break
    for family in sorted({f.strip().strip("'\"").lower() for f in re.findall(r"font-family\s*=\s*[\"']([^\"',]+)", source, re.I)} - allowed_first):
        errors.append(f"SVG font-family {family!r} differs from the shared stack")
    # One object, one icon (visual-design.html 图形：一套线性图标): the same pictogram must not label two
    # different things. Short glyphs (arrows, close, plus, check) are controls, not pictograms, and are skipped.
    if "interface" in ALLOWED or "website" in ALLOWED:
        meaning: dict[str, set[str]] = {}
        # Rows of one list that all carry the same object icon (several meetings, several accounts) name one kind of object.
        row_icons = repeated_row_icons(source)
        for m in re.finditer(r"<svg\b[^>]*>(.*?)</svg>\s*(?:<[^>]+>\s*)*([^<>{}'\"+]{2,16})", source, re.S):
            if m.start() in row_icons:
                continue
            drawing = re.sub(r"\s+", "", m.group(1))
            if len(drawing) < 60 or "<text" in drawing:
                continue
            label = re.sub(r"[\d\s·（）()¥%、，,.]+", "", m.group(2))
            if len(label) >= 2:
                meaning.setdefault(drawing, set()).add(label)
        for drawing, labels in meaning.items():
            distinct = {l for l in labels if not any(l != o and (l in o or o in l) for o in labels)}
            if len(distinct) > 1:
                errors.append(f"one icon labels different things: {sorted(distinct)[:4]}; give each object its own icon")
    # Elements the browser renders in italic by default must be reset to normal when they carry text.
    italic_tags = sorted({t.lower() for t, inner in re.findall(r"<(i|em|cite|var|dfn|address)\b[^>]*>([^<]*)</\1>", source, re.I) if inner.strip()})
    resets = re.findall(r"([^{}]*)\{[^{}]*font-style\s*:\s*normal[^{}]*\}", css, re.I)
    for tag in italic_tags:
        if not any(re.search(rf"(?:^|[\s,>+~(]){tag}\b|(?:^|[\s,])\*(?:[\s,]|$)", sel) for sel in resets):
            errors.append(f"<{tag}> renders italic by default; reset it with font-style: normal")
    # Headings below 22px, table headers and bold tags render 700 by default; the rule is 600 there (visual-design.html 字体与字号).
    # An occurrence that sets its own font-weight inline (pages pasted into editors that strip classes) is already decided.
    bold_tags = sorted({t.lower() for t, attrs, inner in re.findall(r"<(th|strong|b|h3|h4|h5|h6)\b([^>]*)>(.*?)</\1>", source, re.I | re.S)
                        if re.sub(r"<[^>]+>", "", inner).strip()
                        and not re.search(r"""style\s*=\s*["'][^"']*font-weight\s*:""", attrs, re.I)})
    weight_rules = [sel for sel, decls in re.findall(r"([^{}]*)\{([^{}]*)\}", css)
                    if re.search(r"(?:^|;)\s*font-weight\s*:", decls, re.I)
                    or re.search(r"(?:^|;)\s*font\s*:\s*[^;]*\b(?:[1-9]00|bold|normal)\b", decls, re.I)]
    for tag in bold_tags:
        if not any(re.search(rf"(?:^|[\s,>+~(]){tag}\b|(?:^|[\s,])\*(?:[\s,]|$)", sel) for sel in weight_rules):
            warnings.append(f"<{tag}> renders 700 by default; set font-weight: 600 (700 is only for titles and numbers at 22px and larger)")
    radii = {v for v in re.findall(r"border-radius\s*:\s*(\d+)px", css) if v not in {"0", "999", "9999", "50"}}
    if len(radii) > 4:
        warnings.append(f"inconsistent corner radii: {sorted(radii, key=int)} (tags 4, controls 8, cards 12, large cards and dialogs 16)")
    # Templates are skeletons (SKILL.md, WORKFLOW.md): a page that keeps a template's whole style sheet or its default
    # :root palette was not designed for its content.
    page_style = re.sub(r"\s+", "", "\n".join(re.findall(r"<style[^>]*>(.*?)</style>", source, re.S | re.I)))
    for name, template_css in template_styles():
        root_block = re.search(r":root\{[^}]*\}", template_css)
        if page_style and (page_style == template_css or (root_block and root_block.group(0) in page_style)):
            warnings.append(f"page keeps the default style of template {name}: templates are skeletons; design colour, type, "
                            "elevation and graphics for this content from visual-design.html and themes.json")
            break
    if "interface" in ALLOWED:
        for m in re.finditer(r"<svg\b[^>]*\bclass=[\"'][^\"']*\b(?:icon|i|fi|ico)\b[^\"']*[\"'][^>]*>", source, re.I):
            width = re.search(r"\bstroke-width=[\"']([\d.]+)", m.group(0))
            if width:
                icon_strokes.add(width.group(1))
        widths = {float(w) for w in icon_strokes}
        if len(widths) > 1 and not widths <= ICON_STROKES:
            errors.append(f"mixed icon stroke widths {sorted(icon_strokes)}: one icon set, one stroke width (扁平线性风格)")
        for width in sorted(widths - ICON_STROKES):
            warnings.append(f"icon stroke width {width:g}: line icons use 2 on the 24 grid, 1.75 at the least for 16px icons (图形：一套线性图标)")
        # Duotone icons and layered illustrations fill with theme steps or light neutral grey (visual-design.html 插画).
        palette = theme_palette() | {c for v in variables.values() for c in parse_colors(resolve(v))}
        off_theme: list[str] = []
        fills = re.findall(r"<(?:svg|g|path|rect|circle|ellipse|polygon|polyline|line)\b[^>]*?\bfill=[\"']([^\"']+)[\"']", source, re.I)
        fills += [m.group(1) for r in re.finditer(r"([^{}]*)\{([^{}]*)\}", css) if re.search(r"svg|ill|icon|duo|cover|art|\bi\b", r.group(1), re.I)
                  for m in re.finditer(r"(?:^|;)\s*fill\s*:\s*([^;]+)", r.group(2), re.I)]
        for raw in fills:
            for color, alpha in parse_colors_alpha(resolve(raw)):
                hex_value = "#%02X%02X%02X" % color
                if alpha < 1 or color in palette or light_neutral(color) or hex_value in off_theme:
                    continue
                off_theme.append(hex_value)
        if off_theme:
            warnings.append(f"icon or illustration fills outside the theme steps and light neutral grey: {', '.join(off_theme[:6])}"
                            f"{' …' if len(off_theme) > 6 else ''} (take fills from themes.json: primary steps, module accents, neutrals)")
        for m in re.finditer(r"<(linearGradient|radialGradient)\b", source):
            errors.append(f"<{m.group(1)}> in an interface graphic: icons, charts and illustrations are flat")
        for m in re.finditer(r"<img\b[^>]*class=[\"'][^\"']*\b(?:i3|icon|ico)\b", source, re.I):
            errors.append(f"raster icon {m.group(0)[:50]!r}: draw icons as one set of line vectors")
        # Default avatars are cartoon animals or a line icon, never initials, a surname or any other text (Luz-System ui-standards.html, 头像).
        for m in re.finditer(r"class=[\"'][^\"']*\bavatar\b[^\"']*[\"'][^>]*>\s*([^<\s'][^<]{0,5}?)\s*<", source, re.I):
            errors.append(f"text inside avatar: {m.group(1)!r}; use a cartoon avatar or a line icon")
        for m in re.finditer(r"avatar[^\n]{0,80}?(?:\[0\]|charAt\(0\)|slice\(0,\s*[12]\)|substr(?:ing)?\(0,\s*[12]\)|\.(?:name|nick\w*|user\w*)\b[^\n]{0,6}?(?:\[0\]|\.slice))", source, re.I):
            errors.append(f"initials or text used as avatar: {m.group(0)[:60]!r}")
        check_interface_layout(source, resolve, errors, warnings, bold_rules)
    if not re.search(r"<html[^>]+lang=", source, re.I):
        warnings.append("missing lang attribute")
    if not re.search(r"name=[\"']viewport", source, re.I):
        warnings.append("missing viewport meta")


# ---------------------------------------------------------------- Interface layout (--allow interface)
# Density, titles, navigation and table columns (visual-design.html 界面密度与表格对齐; Luz-System admin-design.html).
# Only the static DOM is seen: content a script renders at run time is checked in the browser by the prototype tests.

BODY_TAGS = {"td", "th", "p", "li", "label", "input", "button", "select", "textarea"}
# Auxiliary text may use the 12–13px step: timestamps, units, badges, hints, captions.
AUX_CLASS = re.compile(r"meta|hint|time|unit|badge|caption|aux", re.I)
EMPTY_CELL = {"", "—", "-", "–", "--", "——", "/"}
CONTROL_TAGS = {"button", "a", "input", "select", "textarea", "img", "svg", "progress", "meter"}
NAV_CLASS = re.compile(r"nav|menu|side-?bar|\bside\b", re.I)
ACTIVE_CLASS = re.compile(r"active|current|selected", re.I)
STATE_PSEUDO = re.compile(r":(?:hover|focus|active|disabled|checked|visited|focus-within|focus-visible|placeholder|invalid|valid|"
                          r"empty|target|indeterminate|read-only|placeholder-shown|-webkit-[\w-]+|-moz-[\w-]+)\b|::|:before|:after|:marker", re.I)


class Node:
    __slots__ = ("tag", "attrs", "children", "parent", "hidden", "classes", "pos")

    def __init__(self, tag: str, attrs: dict, parent: "Node | None") -> None:
        self.tag, self.attrs, self.parent, self.children = tag, attrs, parent, []
        self.pos: tuple[int, int] | None = None
        self.classes = set((attrs.get("class") or "").split())
        self.hidden = bool(parent and parent.hidden) or tag in {"script", "style", "template", "head", "noscript"} \
            or "hidden" in attrs or (attrs.get("aria-hidden") or "").lower() == "true" \
            or bool(re.search(r"display\s*:\s*none", attrs.get("style") or "", re.I))

    def elements(self):
        for child in self.children:
            if isinstance(child, Node):
                yield child
                yield from child.elements()

    def text(self) -> str:
        parts = []
        for child in self.children:
            if isinstance(child, str):
                parts.append(child)
            elif child.tag not in {"script", "style", "template"}:
                parts.append(child.text())
        return " ".join(" ".join(parts).split())

    def own_text(self) -> str:
        return " ".join(" ".join(c for c in self.children if isinstance(c, str)).split())

    def ancestors(self):
        node = self.parent
        while node is not None:
            yield node
            node = node.parent

    def cls(self) -> str:
        return self.attrs.get("class") or ""


class TreeBuilder(HTMLParser):
    """A small DOM: enough structure for selector matching, tables and navigation groups."""

    VOID = Outline.VOID | {"area", "base", "embed", "param", "track", "stop"}
    # Opening one of these closes an open element of the listed kinds first (HTML optional end tags).
    AUTO_CLOSE = {"li": ({"li"}, {"ul", "ol", "menu"}), "td": ({"td", "th"}, {"tr", "table"}), "th": ({"td", "th"}, {"tr", "table"}),
                  "tr": ({"tr", "td", "th"}, {"table", "tbody", "thead", "tfoot"}), "p": ({"p"}, set()),
                  "option": ({"option"}, {"select"}), "dt": ({"dt", "dd"}, {"dl"}), "dd": ({"dt", "dd"}, {"dl"})}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.root = Node("#root", {}, None)
        self.current = self.root

    def handle_starttag(self, tag, attrs):
        if tag in self.AUTO_CLOSE:
            closes, stops = self.AUTO_CLOSE[tag]
            node = self.current
            while node is not self.root and node.tag not in stops:
                if node.tag in closes:
                    self.current = node.parent
                    break
                node = node.parent
        node = Node(tag, {k: (v if v is not None else "") for k, v in attrs}, self.current)
        node.pos = self.getpos()
        self.current.children.append(node)
        if tag not in self.VOID:
            self.current = node

    def handle_startendtag(self, tag, attrs):
        node = Node(tag, {k: (v if v is not None else "") for k, v in attrs}, self.current)
        node.pos = self.getpos()
        self.current.children.append(node)

    def handle_endtag(self, tag):
        node = self.current
        while node is not self.root and node.tag != tag:
            node = node.parent
        if node is not self.root:
            self.current = node.parent

    def handle_data(self, data):
        if data.strip():
            self.current.children.append(data)


def repeated_row_icons(source: str) -> set[int]:
    """Source offsets of <svg> icons that mark the kind of object in a repeated list row: the svg sits in one of two or more
    sibling rows of the same tag and class, every one of those rows carries the same drawing, and every row holds at
    least two pieces of text one of which has a digit (an instance's name plus its time, date, count or amount: each
    meeting in a meeting list, each account in an account list). Entry grids (an icon and one label), rows that mix
    drawings and feature lists that hang one generic icon before different titles are not exempt, so different objects
    sharing one pictogram are still reported (visual-design.html 图形：一套线性图标)."""
    builder = TreeBuilder()
    builder.feed(source)
    line_starts = [0] + [m.end() for m in re.finditer(r"\n", source)]

    def drawing(svg: Node) -> str:
        return "".join(f"<{n.tag}{sorted(n.attrs.items())}>" for n in svg.elements())

    def text_pieces(node: Node) -> int:
        if node.tag == "svg":
            return 0
        return sum(1 if isinstance(c, str) and c.strip() else text_pieces(c) if isinstance(c, Node) else 0 for c in node.children)

    exempt: set[int] = set()
    for svg in (n for n in builder.root.elements() if n.tag == "svg" and n.pos):
        sig = drawing(svg)
        if not sig:
            continue
        row = svg.parent
        for _ in range(5):
            parent = row.parent if row is not None else None
            if parent is None or row.tag == "#root" or parent.tag == "#root":
                break
            rows = [c for c in parent.children if isinstance(c, Node) and c.tag == row.tag and c.cls() == row.cls()]
            if len(rows) >= 2:
                if all(text_pieces(r) >= 2 and re.search(r"\d", r.text()) and any(drawing(d) == sig for d in r.elements() if d.tag == "svg")
                       for r in rows):
                    line, col = svg.pos
                    exempt.add(line_starts[line - 1] + col)
                break
            row = parent
    return exempt


def css_rules(css: str) -> list[tuple[str, str]]:
    """Top-level (selector, declarations) pairs; rules inside @media and other at-rules are skipped, so a mobile
    override does not decide the desktop cascade."""
    css = re.sub(r"/\*.*?\*/", "", css, flags=re.S)
    rules, depth, start, selector = [], 0, 0, ""
    for i, ch in enumerate(css):
        if ch == "{":
            if depth == 0:
                selector = css[start:i].strip()
            depth += 1
            start = i + 1
        elif ch == "}":
            depth = max(depth - 1, 0)
            if depth == 0 and not selector.startswith("@"):
                rules.append((selector, css[start:i]))
            start = i + 1
        elif ch == ";" and depth == 0:
            start = i + 1
    return rules


def parse_compound(text: str):
    m = re.match(r"([a-zA-Z][\w-]*|\*)?", text)
    tag = (m.group(1) or "*").lower()
    rest = text[m.end():]
    ids = re.findall(r"#([\w-]+)", rest)
    classes = re.findall(r"\.([\w-]+)", rest)
    attrs = re.findall(r"\[\s*([\w-]+)\s*(?:([~*^$|]?=)\s*[\"']?([^\"'\]]*)[\"']?)?\s*\]", rest)
    pseudo = len(re.findall(r":[\w-]+", re.sub(r"\[[^\]]*\]", "", rest)))
    return tag, ids, classes, attrs, (len(ids), len(classes) + len(attrs) + pseudo, 0 if tag == "*" else 1)


def parse_selector(selector: str):
    """[(combinator, compound)] read right to left; None for selectors that depend on state, pseudo-elements or siblings."""
    selector = selector.strip()
    if not selector or STATE_PSEUDO.search(selector) or re.search(r"[+~](?![^\[]*\])", selector):
        return None
    selector = re.sub(r":(?:not|is|where|has|nth-[\w-]+)\([^)]*\)", "", selector)
    selector = re.sub(r":root\b", "html", selector)
    selector = re.sub(r"\s*>\s*", " > ", selector)
    parts, combinator = [], " "
    for token in selector.split():
        if token == ">":
            combinator = ">"
            continue
        parts.append((combinator, parse_compound(token)))
        combinator = " "
    return parts[::-1] if parts else None


def match_compound(node: Node, compound) -> bool:
    tag, ids, classes, attrs, _ = compound
    if tag != "*" and node.tag != tag:
        return False
    if any(node.attrs.get("id") != i for i in ids) or not set(classes) <= node.classes:
        return False
    for name, op, value in attrs:
        if name not in node.attrs:
            return False
        actual = node.attrs[name]
        if op == "=" and actual != value or op == "~=" and value not in actual.split() or op == "*=" and value not in actual \
                or op == "^=" and not actual.startswith(value) or op == "$=" and not actual.endswith(value):
            return False
    return True


def match_selector(node: Node, parts, index: int = 0) -> bool:
    if not match_compound(node, parts[index][1]):
        return False
    if index + 1 == len(parts):
        return True
    combinator = parts[index][0]
    for ancestor in node.ancestors():
        if ancestor.tag == "#root":
            return False
        if match_selector(ancestor, parts, index + 1):
            return True
        if combinator == ">":
            return False
    return False


def length_px(value: str, parent: float) -> float | None:
    m = re.search(r"(-?(?:\d+(?:\.\d+)?|\.\d+))\s*(px|rem|em|pt|%)?", value)
    if not m or "calc" in value or "clamp" in value or "var(" in value:
        return None
    number, unit = float(m.group(1)), (m.group(2) or "px").lower()
    return {"px": number, "rem": number * 16, "em": number * parent, "pt": number * 4 / 3, "%": number * parent / 100}[unit]


def declarations(body: str, resolve) -> dict[str, str]:
    found = {}
    for name, value in re.findall(r"(?:^|;)\s*([\w-]+)\s*:\s*([^;]+)", body):
        found[name.lower()] = resolve(value.replace("!important", "").strip())
    return found


def font_size_value(decls: dict[str, str]) -> str | None:
    if "font-size" in decls:
        return decls["font-size"]
    if "font" in decls:
        m = re.search(r"((?:\d+(?:\.\d+)?|\.\d+)(?:px|rem|em|pt|%))", decls["font"])
        return m.group(1) if m else None
    return None


def padding_sides(decls: dict[str, str]) -> tuple[str | None, str | None]:
    left = right = None
    for name in ("padding", "padding-inline"):
        if name in decls:
            values = decls[name].split()
            if name == "padding":
                left = right = values[1] if len(values) >= 2 else values[0]
                if len(values) == 4:
                    left = values[3]
            else:
                left, right = values[0], values[-1]
    left = decls.get("padding-left", decls.get("padding-inline-start", left))
    right = decls.get("padding-right", decls.get("padding-inline-end", right))
    return left, right


def check_interface_layout(source: str, resolve, errors: list[str], warnings: list[str],
                           bold_rules: list[tuple[str, str]] | None = None) -> None:
    builder = TreeBuilder()
    builder.feed(source)
    nodes = [n for n in builder.root.elements() if not n.hidden and not any(a.tag == "svg" for a in n.ancestors())]
    by_key: dict[str, list[Node]] = {}
    for n in nodes:
        for key in [n.tag] + ["." + c for c in n.classes] + (["#" + n.attrs["id"]] if n.attrs.get("id") else []):
            by_key.setdefault(key, []).append(n)

    # Cascade: for each element, the winning font-size and padding by specificity, then source order; inline style wins.
    css = "\n".join(re.findall(r"<style[^>]*>(.*?)</style>", source, re.S | re.I))
    won: dict[int, dict[str, tuple]] = {}
    for order, (group, body) in enumerate(css_rules(css)):
        decls = declarations(body, resolve)
        wanted = {}
        size = font_size_value(decls)
        if size:
            wanted["size"] = size
        left, right = padding_sides(decls)
        if left is not None:
            wanted["pad-left"] = left
        if right is not None:
            wanted["pad-right"] = right
        if not wanted:
            continue
        for selector in group.split(","):
            parts = parse_selector(selector)
            if not parts:
                continue
            tag, ids, classes, _, spec = parts[0][1]
            key = "#" + ids[0] if ids else "." + classes[0] if classes else tag
            candidates = nodes if key == "*" else by_key.get(key, [])
            for n in candidates:
                if match_selector(n, parts):
                    slot = won.setdefault(id(n), {})
                    for prop, value in wanted.items():
                        if prop not in slot or (spec, order) >= slot[prop][0]:
                            slot[prop] = ((spec, order), value)

    sizes: dict[int, float] = {}

    def size_of(n: Node) -> float:
        if id(n) in sizes:
            return sizes[id(n)]
        parent = size_of(n.parent) if n.parent is not None and n.parent.tag != "#root" else 16.0
        inline = declarations(n.attrs.get("style") or "", resolve)
        value = font_size_value(inline) or won.get(id(n), {}).get("size", (None, None))[1]
        px = length_px(value, parent) if value else None
        if value and value.strip() in {"inherit", "unset"}:
            px = None
        sizes[id(n)] = parent if px is None else px
        return sizes[id(n)]

    # 1. Body text 14px: cells, paragraphs, list items, labels and controls. Below 12px is already an error above.
    small: dict[tuple, list[str]] = {}
    for n in nodes:
        if n.tag not in BODY_TAGS:
            continue
        if n.tag == "input" and (n.attrs.get("type") or "text").lower() in {"hidden", "checkbox", "radio", "range", "color", "file", "image"}:
            continue
        text = n.own_text() if n.tag not in {"input", "select", "textarea"} else (n.attrs.get("value") or n.attrs.get("placeholder") or n.tag)
        if not text:
            continue
        px = size_of(n)
        if 12 <= px < 14:
            chain = [n] + list(n.ancestors())[:2]
            if any(AUX_CLASS.search(a.cls()) for a in chain) or n.tag in {"small", "time"}:
                continue
            small.setdefault((n.tag, n.cls(), round(px, 1)), []).append(text)
    for (tag, cls, px), texts in small.items():
        label = f"<{tag}{' class=' + repr(cls) if cls else ''}>"
        errors.append(f"interface body text {px:g}px in {label} ({len(texts)}x, e.g. {texts[0][:16]!r}): body, tables and forms use 14px; "
                      "12–13px only for elements marked meta, hint, time, unit, badge, caption or aux")

    def in_nav(n: Node) -> bool:
        return any(a.tag in {"nav", "aside"} or (a.attrs.get("role") or "") in {"navigation", "menu", "menubar", "tree"}
                   or NAV_CLASS.search(a.cls()) for a in [n] + list(n.ancestors()))

    def norm(text: str) -> str:
        return re.sub(r"\s+", "", text)

    # 2. The page title repeats the highlighted navigation item.
    active = [n for n in nodes if in_nav(n) and (ACTIVE_CLASS.search(n.cls()) or (n.attrs.get("aria-current") or "false").lower() != "false"
                                                 and "aria-current" in n.attrs)]
    active_text = {norm(n.text()): n for n in active if norm(n.text())}
    titles = [n for n in nodes if not in_nav(n) and (
        n.tag == "h1" or re.search(r"page[-_]?(?:title|head|header|name)", n.cls(), re.I)
        or (n.tag in {"h2", "h3", "h4"} and any(re.search(r"page[-_]?(?:title|head)|(?:^|[-_\s])title\b", a.cls(), re.I) for a in [n] + list(n.ancestors())[:2])))]
    reported: set[str] = set()
    for t in titles:
        text = norm(t.text())
        if text in active_text and text not in reported:
            reported.add(text)
            errors.append(f"page title {t.text()!r} repeats the highlighted navigation item: drop the title and use the header row for actions and filters")

    # 3. A navigation group whose label equals its only item.
    def clickables(scope: list[Node]) -> list[Node]:
        found = [n for n in scope if n.tag in {"a", "button"} or (n.attrs.get("role") or "") in {"menuitem", "link", "tab", "treeitem"}]
        return found or [n for n in scope if n.tag == "li"]

    for label in nodes:
        if not in_nav(label) or not label.text() or label.tag in {"a", "button", "li"} or \
                any(a.tag in {"a", "button"} for a in label.ancestors()):
            continue
        if not (label.tag in {"h2", "h3", "h4", "h5", "h6", "dt"} or re.search(r"label|title|head|caption|group-?name|\bmg\b", label.cls(), re.I)):
            continue
        siblings = [c for c in label.parent.children if isinstance(c, Node)] if label.parent else []
        after = siblings[siblings.index(label) + 1:] if label in siblings else []
        items = clickables([x for s in after[:1] for x in [s] + list(s.elements())]) if after else []
        if not items and label.parent is not None:
            items = clickables([x for x in label.parent.elements() if x is not label and label not in x.ancestors()])
        if len(items) == 1 and norm(items[0].text()) == norm(label.text()):
            errors.append(f"navigation group {label.text()!r} holds only an item of the same name: remove the group label")

    # 4–6. Table columns: all empty, all identical; long text without 16px side padding.
    padding_reported: set[tuple] = set()
    for table in (n for n in nodes if n.tag == "table"):
        rows = [r for r in table.elements() if r.tag == "tr" and not r.hidden and next((a for a in r.ancestors() if a.tag == "table"), None) is table]
        header = next((r for r in rows if any(c.tag == "th" for c in r.children if isinstance(c, Node))), None)
        names = [c.text() for c in header.children if isinstance(c, Node) and c.tag in {"th", "td"}] if header else []
        # Lay the body rows onto a grid so a cell merged down several rows (rowspan) counts once, in its own column.
        data, carry = [], {}
        for r in rows:
            cells = [c for c in r.children if isinstance(c, Node) and c.tag in {"td", "th"}]
            if r is header or not cells or any(a.tag == "thead" for a in r.ancestors()) or all(c.tag == "th" for c in cells) \
                    or any(c.attrs.get("colspan", "1") not in {"", "1"} for c in cells):
                carry = {}
                continue
            slots, k, queue = [], 0, list(cells)
            while queue or any(key >= k for key in carry):
                if k in carry:
                    slots.append(None)
                    carry[k] -= 1
                    if carry[k] == 0:
                        del carry[k]
                else:
                    cell = queue.pop(0)
                    span = int(cell.attrs.get("rowspan") or 1) if (cell.attrs.get("rowspan") or "1").isdigit() else 1
                    if span > 1:
                        carry[k] = span - 1
                    slots.append(cell)
                k += 1
            data.append(slots)
        width = len(data[0]) if data else 0
        if len(data) < 3 or any(len(r) != width for r in data):
            continue
        for k in range(width):
            column = [r[k] for r in data if r[k] is not None]
            if len(column) < 3:
                continue
            name = names[k] if k < len(names) and len(names) == width else f"第 {k + 1} 列"
            has_control = any(x.tag in CONTROL_TAGS for c in column for x in c.elements())
            texts = [c.text() for c in column]
            if not has_control and all(t in EMPTY_CELL for t in texts):
                errors.append(f"table column {name!r} is empty in every row: hide it by default (column settings)")
            elif len(column) >= 5 and not has_control and len(set(texts)) == 1 and texts[0] not in EMPTY_CELL and not re.search(r"操作", name):
                warnings.append(f"table column {name!r} holds {texts[0][:16]!r} in all {len(column)} rows: state it once in the filter summary or header")
            for c in column:
                if len(c.text()) < 12 or any(x.tag in CONTROL_TAGS for x in c.elements()):
                    continue
                slot = won.get(id(c), {})
                values = [length_px(slot[p][1], size_of(c)) for p in ("pad-left", "pad-right") if p in slot]
                values = [v for v in values if v is not None]
                key = (table.cls(), c.cls(), tuple(values))
                if values and min(values) < 16 and key not in padding_reported:
                    padding_reported.add(key)
                    warnings.append(f"long-text cell in column {name!r} has {min(values):g}px side padding: data cells keep at least 16px")
                break

    # 700 only at 22px and larger (visual-design.html 字体与字号): judge each 700 rule by its own size, else by the computed
    # size of the elements it matches; a 700 rule that matches nothing and states no size cannot be judged and is refused.
    for selector, body in bold_rules or []:
        own = font_size_value(declarations(body, resolve))
        own_px = length_px(own, 16.0) if own else None
        if own_px is not None:
            if own_px < BOLD_MIN_PX:
                errors.append(f"font weight 700 at {own_px:g}px in {selector[:40]!r}: 700 is only for page titles and large numbers at {BOLD_MIN_PX:g}px and larger; use 600")
            continue
        matched: list[Node] = []
        for part in selector.split(","):
            plain = re.sub(r"::?[\w-]+(?:\([^)]*\))?", "", part).strip()
            parsed = parse_selector(plain) if plain else None
            if parsed:
                matched += [n for n in nodes if match_selector(n, parsed)]
        if matched:
            small_sizes = sorted({round(size_of(n), 1) for n in matched if size_of(n) < BOLD_MIN_PX})
            if small_sizes:
                errors.append(f"font weight 700 in {selector[:40]!r} on {', '.join(f'{v:g}px' for v in small_sizes)} text: "
                              f"700 is only for {BOLD_MIN_PX:g}px and larger; use 600")
        else:
            errors.append(f"font weight 700 in {selector[:40]!r} without a font size: 700 is only for page titles and large numbers "
                          f"at {BOLD_MIN_PX:g}px and larger; state the size in the same rule or use 600")
    for n in nodes:
        inline = declarations(n.attrs.get("style") or "", resolve)
        if inline.get("font-weight", "").strip().lower() in BOLD_WEIGHTS and size_of(n) < BOLD_MIN_PX:
            errors.append(f"inline font weight 700 on <{n.tag}> at {size_of(n):g}px: 700 is only for {BOLD_MIN_PX:g}px and larger")

    # One brand card per screen (visual-design.html 配色, 品牌主卡): a screen is the nearest element carrying data-page-id
    # or a phone frame class; without one, the whole page is one screen.
    screens: dict[int, list[Node]] = {}
    for n in builder.root.elements():
        if BRAND_CARD_CLASS not in n.classes:
            continue
        scope = next((a for a in n.ancestors() if a.tag != "#root" and ("data-page-id" in a.attrs or PHONE_CLASS.search(a.cls()))), None)
        screens.setdefault(id(scope) if scope is not None else 0, []).append(n)
    for cards in screens.values():
        if len(cards) > 1:
            where = next((a.attrs.get("data-page-id") for a in cards[0].ancestors() if a.attrs.get("data-page-id")), "the page")
            errors.append(f"{len(cards)} brand cards on one screen ({where}): a screen carries at most one solid brand card")

    check_pagination_and_brand(builder.root, warnings)


# 7–9. List pagination and the brand mark (visual-design.html 分页 and 网页图标与品牌标识; Luz-System admin-design.html 分页).
PAGER_CLASS = re.compile(r"pagination|pager", re.I)
LIST_CLASS = re.compile(r"list|table", re.I)
PAGE_SIZE_TEXT = re.compile(r"(\d+)\s*条\s*/\s*页")
TOTAL_TEXT = re.compile(r"共\s*[\d,]+\s*条")
SKIP_TABLE = re.compile(r"skeleton|loading|export|modal|dialog|detail|popover|tooltip|legend", re.I)
SHELL_CLASS = re.compile(r"(?:^|[-_\s])(?:header|topbar|top-bar|navbar|nav-bar|appbar|app-bar|masthead|sidebar|side-bar|"
                         r"site-head|site-header|admin|console)(?:$|[-_\s])", re.I)
PHONE_CLASS = re.compile(r"(?:^|[-_\s])(?:mp|phone|mobile|miniprogram|mini-program|wx|device|screen-phone)(?:$|[-_\s])", re.I)
BRAND_CLASS = re.compile(r"logo|brand(?!-card)", re.I)
DEFAULT_PAGE_SIZE = 20


def default_page_size(scope: Node) -> int | None:
    """The page size a pagination control shows before the user changes it."""
    for n in [scope] + list(scope.elements()):
        for attr in ("data-page-size", "data-default-size", "data-size"):
            if (n.attrs.get(attr) or "").strip().isdigit():
                return int(n.attrs[attr].strip())
    for select in (n for n in scope.elements() if n.tag == "select"):
        options = [o for o in select.elements() if o.tag == "option"]
        sized = [o for o in options if PAGE_SIZE_TEXT.search(o.text()) or (o.attrs.get("value") or "").strip() in {"10", "20", "50", "100"}]
        if not sized:
            continue
        chosen = next((o for o in sized if "selected" in o.attrs), sized[0])
        value = (chosen.attrs.get("value") or "").strip()
        if value.isdigit():
            return int(value)
        match = PAGE_SIZE_TEXT.search(chosen.text())
        if match:
            return int(match.group(1))
    match = PAGE_SIZE_TEXT.search(scope.text())
    return int(match.group(1)) if match else None


def check_pagination_and_brand(root: Node, warnings: list[str]) -> None:
    elements = [n for n in root.elements() if not any(a.tag in {"template", "script"} for a in n.ancestors())]

    def chain(n: Node) -> list[Node]:
        return [n] + list(n.ancestors())

    # 7. A list table (10+ data rows, or a table marked as a list) without a pagination control.
    def data_rows(table: Node) -> int:
        return sum(1 for r in table.elements() if r.tag == "tr" and next((a for a in r.ancestors() if a.tag == "table"), None) is table
                   and not any(a.tag == "thead" for a in r.ancestors()) and any(isinstance(c, Node) and c.tag == "td" for c in r.children))

    lists = []
    for table in (n for n in elements if n.tag == "table" and not n.hidden):
        if any((a.attrs.get("data-pagination") or "").lower() in {"none", "off", "false"} or "data-state-kind" in a.attrs
               or "data-export-preview" in a.attrs or SKIP_TABLE.search(a.cls()) for a in chain(table)):
            continue
        marked = LIST_CLASS.search(table.cls()) or (table.parent is not None and LIST_CLASS.search(table.parent.cls()))
        if marked or data_rows(table) >= 10:
            lists.append(table)
    pagers = [n for n in elements if PAGER_CLASS.search(n.cls())]
    page_text = root.text()
    text_pager = bool(TOTAL_TEXT.search(page_text) and PAGE_SIZE_TEXT.search(page_text))
    if lists and not pagers and not text_pager:
        sample = lists[0].cls() or (lists[0].parent.cls() if lists[0].parent is not None else "")
        warnings.append(f"list table without pagination ({len(lists)} list table(s){', e.g. class ' + repr(sample) if sample else ''}): "
                        "add 共 N 条, a 10/20/50/100 条/页 selector and page numbers below the table, default 20; "
                        'mark a list that can never exceed one page with data-pagination="none"')
    # 8. The default page size is 20.
    sizes = {default_page_size(p) for p in pagers} if pagers else ({default_page_size(root)} if text_pager else set())
    for size in sorted(s for s in sizes if s is not None and s != DEFAULT_PAGE_SIZE):
        warnings.append(f"pagination defaults to {size} 条/页: default to {DEFAULT_PAGE_SIZE} (options 10/20/50/100) and remember the user's choice")

    # 9. An admin or website shell (header, side bar, top bar) carries a graphic brand mark: icon + system name.
    shells = [n for n in elements if (n.tag in {"header", "aside"} or SHELL_CLASS.search(n.cls()))
              and not any(PHONE_CLASS.search(a.cls()) for a in chain(n))]
    if shells:
        holders = [n for n in elements if n.tag == "header" or BRAND_CLASS.search(n.cls())]
        graphic = any(h.tag in {"svg", "img"} or any(d.tag in {"svg", "img"} for d in h.elements()) for h in holders)
        if not graphic:
            warnings.append("admin or website header without a graphic brand mark: put the tab-icon pictogram (svg or img) "
                            "beside the system name in a header, logo or brand element, never a letter tile")


# ---------------------------------------------------------------- XLSX

def check_xlsx(path: Path, errors: list[str], warnings: list[str]) -> None:
    try:
        import openpyxl
    except ImportError:
        errors.append("openpyxl is required: py -m pip install openpyxl")
        return
    book = openpyxl.load_workbook(path)
    sheets = book.worksheets
    if len(sheets) > 1 and "multi-sheet" not in ALLOWED:
        errors.append(f"{len(sheets)} worksheets; only one is allowed unless the user asked for more ({', '.join(s.title for s in sheets)})")
    fonts: set[str] = set()
    sizes: set[float] = set()
    for sheet in sheets:
        where = f"[{sheet.title}]"
        if re.fullmatch(r"(?:Sheet|工作表)\d*", sheet.title, re.I):
            errors.append(f"{where} default sheet name; name it after the content")
        if sheet.sheet_state != "visible":
            errors.append(f"{where} hidden worksheet")
        if getattr(sheet, "_images", None) or getattr(sheet, "_charts", None):
            if "chart" not in ALLOWED:
                errors.append(f"{where} images or charts the user did not ask for (pass --allow chart if requested)")
        rows = [[c for c in row if c.value not in (None, "")] for row in sheet.iter_rows()]
        while rows and not rows[-1]:
            rows.pop()
        if not rows or not rows[0]:
            errors.append(f"{where} row 1 must hold the title")
            continue
        if len(rows[0]) != 1 or rows[0][0].column != 1:
            errors.append(f"{where} row 1 must contain only the title in A1")
        header_index = next((i for i, r in enumerate(rows[1:], 1) if len(r) >= 2), None)
        if header_index is None:
            errors.append(f"{where} no header row found")
            continue
        for i in range(1, header_index):
            if rows[i]:
                errors.append(f"{where} row {i + 1} sits between the title and the header: {str(rows[i][0].value)[:30]!r}")
        header = rows[header_index]
        header_values = [str(c.value).strip() for c in header]
        # Exported data speaks business Chinese, not database fields (file-design.html, 导出文件).
        fields = [v for v in header_values if re.fullmatch(r"[a-z][a-z0-9]*(?:_[a-z0-9]+)+|[a-z][a-z0-9]*(?:[A-Z][a-z0-9]*)+|[a-z][a-z0-9]{1,}", v)]
        if fields:
            warnings.append(f"{where} header holds field names {fields[:4]}: translate them into Chinese business terms")
        last_col = max(c.column for c in header)
        first_col = min(c.column for c in header)
        # Everything must sit inside the data rectangle: title, header, data, section titles, totals.
        for i in range(header_index + 1, len(rows)):
            row = rows[i]
            if not row:
                continue
            outside = [c for c in row if c.column > last_col or c.column < first_col]
            for cell in outside:
                errors.append(f"{where}{cell.coordinate} content outside the data table: {str(cell.value)[:30]!r}")
            inside = [c for c in row if c not in outside]
            # A row that continues a vertically merged group cell is a data row, not a lone note.
            in_group = any(m.min_row > 1 and m.min_row < inside[0].row <= m.max_row for m in sheet.merged_cells.ranges) if inside else False
            if len(inside) == 1 and isinstance(inside[0].value, str) and len(header) > 1 and not in_group:
                following = next((r for r in rows[i + 1:] if r), None)
                is_section = following is not None and [str(c.value).strip() for c in following] == header_values
                if not is_section:
                    errors.append(f"{where}{inside[0].coordinate} single text cell inside or below the table "
                                  f"(note, legend, unit or conclusion?): {inside[0].value[:30]!r}")
        title_size = rows[0][0].font.sz or 11
        for row in rows:
            for cell in row:
                if isinstance(cell.value, str):
                    check_text(cell.value, f"{where}{cell.coordinate}", errors)
                fonts.add((cell.font.name or "").lower())
                size = float(cell.font.sz or 11)
                sizes.add(size)
                if size < 12:
                    errors.append(f"{where}{cell.coordinate} font {size:g}pt is below 12pt")
                if size > title_size:
                    errors.append(f"{where}{cell.coordinate} is larger than the title")
                fill = cell.fill
                if fill is not None and fill.fill_type == "solid":
                    rgb = getattr(fill.fgColor, "rgb", None)
                    if isinstance(rgb, str) and re.fullmatch(r"[0-9A-Fa-f]{8}", rgb) and luminance(hex_rgb(rgb)) < 0.6:
                        errors.append(f"{where}{cell.coordinate} dark fill #{rgb[2:]}")
                color = getattr(cell.font.color, "rgb", None) if cell.font.color is not None else None
                if isinstance(color, str) and re.fullmatch(r"[0-9A-Fa-f]{8}", color):
                    contrast = 1.05 / (luminance(hex_rgb(color)) + 0.05)
                    if contrast < 4.5:
                        errors.append(f"{where}{cell.coordinate} font color too light #{color[2:]}")
        if not header[0].font.b:
            warnings.append(f"{where} header row is not bold")
        data_rows = len(rows) - header_index - 1
        if data_rows > 20 and not sheet.freeze_panes:
            errors.append(f"{where} freeze the header row")
        has_vertical_merges = any(m.min_row > 1 and m.max_row > m.min_row for m in sheet.merged_cells.ranges)
        if data_rows > 20 and not sheet.auto_filter.ref and not has_vertical_merges:
            warnings.append(f"{where} no filter on a table longer than 20 rows")
        if not any(d.width not in (None, 13.0) for d in sheet.column_dimensions.values()):
            errors.append(f"{where} column widths were not set")
        # Merging (file-design.html, 表格文件): the title spans the table; inside the data a cell may only merge
        # vertically to group equal consecutive values of one column; horizontal layout merges are not allowed
        # except a label spanning the merged group columns on a total row.
        header_row = header[0].row
        merged_cover: set[tuple[int, int]] = set()
        for merged in sheet.merged_cells.ranges:
            for r in range(merged.min_row, merged.max_row + 1):
                for c in range(merged.min_col, merged.max_col + 1):
                    merged_cover.add((r, c))
            if merged.min_row == 1:
                continue
            if merged.min_col != merged.max_col and merged.min_row == merged.max_row and merged.min_row > header_row:
                label = sheet.cell(merged.min_row, merged.min_col).value
                # The row's first-column label may itself be a vertical merge (合计 over two summary rows).
                row_label = sheet.cell(merged.min_row, first_col).value
                for other in sheet.merged_cells.ranges:
                    if other.min_col <= first_col <= other.max_col and other.min_row <= merged.min_row <= other.max_row:
                        row_label = sheet.cell(other.min_row, other.min_col).value
                if not any(isinstance(v, str) and re.search(r"合计|总计|小计", v) for v in (label, row_label)):
                    errors.append(f"{where} horizontal merge in the data area: {merged.coord}")
            elif merged.min_col != merged.max_col:
                errors.append(f"{where} block merge in the data area: {merged.coord}")
        # No sequence-number column: the data identifies itself.
        for cell in header:
            if re.fullmatch(r"\s*(?:序号|行号|序|No\.?|#)\s*", str(cell.value), re.I):
                errors.append(f"{where}{cell.coordinate} sequence-number column {cell.value!r}; write the data itself")
        data_cells = [r for r in rows[header_index + 1:] if r]
        first_col_values = [r[0].value for r in data_cells if r[0].column == first_col]
        if len(first_col_values) >= 3 and all(isinstance(v, int) for v in first_col_values) \
                and first_col_values == list(range(first_col_values[0], first_col_values[0] + len(first_col_values))):
            errors.append(f"{where} first column is a running number; remove the sequence column")
        # Same text repeated down a grouping column must be merged, not written on every row.
        grouping = True
        for col in range(first_col, last_col + 1):
            values = [(r, sheet.cell(r, col).value) for r in range(header_row + 1, sheet.max_row + 1)]
            filled = [v for _, v in values if v not in (None, "")]
            if not filled or sum(isinstance(v, str) for v in filled) * 2 < len(filled):
                continue  # numbers and sequence columns neither group nor break grouping
            run_found = any(
                isinstance(v, str) and v.strip() not in ("", "—", "-") and v == values[i + 1][1]
                and (r, col) not in merged_cover and (values[i + 1][0], col) not in merged_cover
                for i, (r, v) in enumerate(values[:-1])
            )
            is_group_col = any((r, col) in merged_cover for r, _ in values) or run_found
            if run_found and grouping:
                errors.append(f"{where} column {sheet.cell(header_row, col).value!r} repeats the same text on consecutive rows; merge those cells vertically")
            if not is_group_col:
                grouping = False
        header_fill = header[0].fill
        if header_fill is None or header_fill.fill_type != "solid":
            warnings.append(f"{where} header row has no light fill")
        for row in sheet.iter_rows():
            for cell in row:
                if cell.comment:
                    errors.append(f"{where}{cell.coordinate} cell comment")
        hidden_cols = [k for k, d in sheet.column_dimensions.items() if d.hidden]
        hidden_rows = [k for k, d in sheet.row_dimensions.items() if d.hidden]
        if hidden_cols or hidden_rows:
            errors.append(f"{where} hidden rows or columns: {hidden_cols + hidden_rows}")
        if not sheet.print_area:
            warnings.append(f"{where} print area not set")
    if len(fonts) > 1:
        errors.append(f"more than one font family: {sorted(fonts)}")
    if len(sizes) > 4:
        errors.append(f"more than four font sizes: {sorted(sizes)}")


# ---------------------------------------------------------------- DOCX

def check_docx(path: Path, errors: list[str], warnings: list[str]) -> None:
    try:
        import docx
        from docx.oxml.ns import qn
    except ImportError:
        errors.append("python-docx is required: py -m pip install python-docx")
        return
    document = docx.Document(str(path))
    paragraphs = [p for p in document.paragraphs if p.text.strip()]
    if not paragraphs:
        errors.append("document has no text")
        return

    def is_heading(paragraph) -> bool:
        name = (paragraph.style.name or "").lower()
        return "heading" in name or "title" in name or "标题" in name

    first = paragraphs[0]
    if not is_heading(first):
        sizes = [r.font.size.pt for r in first.runs if r.font.size]
        body = [r.font.size.pt for p in paragraphs[1:6] for r in p.runs if r.font.size]
        if not sizes or (body and max(sizes) <= max(body)):
            errors.append(f"first paragraph is not a title: {first.text[:40]!r}")
    if len(paragraphs) > 1 and not re.search(r"合同|协议", first.text) and not is_heading(paragraphs[1]) and not all(r.bold for r in paragraphs[1].runs if r.text.strip()):
        errors.append(f"text between the title and the first heading: {paragraphs[1].text[:40]!r}")
    families: set[str] = set()
    for paragraph in paragraphs:
        check_text(paragraph.text, "paragraph", errors)
        if is_heading(paragraph):
            check_heading(paragraph.text, "heading", errors)
        for run in paragraph.runs:
            if run.font.size and run.font.size.pt < 10.5:
                errors.append(f"font {run.font.size.pt:g}pt below 10.5pt (五号): {run.text[:20]!r}")
            r_fonts = run._element.rPr.rFonts if run._element.rPr is not None else None
            east = r_fonts.get(qn("w:eastAsia")) if r_fonts is not None else None
            for name in (run.font.name, east):
                if name:
                    families.add(name.lower())
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                check_text(cell.text, "table cell", errors)
    if len(families) > 2:
        errors.append(f"more than one font family: {sorted(families)}")
    # Document types (file-design.html, 按文档类型).
    full_text = "\n".join([p.text for p in document.paragraphs] + [c.text for t in document.tables for r in t.rows for c in r.cells])
    for m in re.finditer(r"【[^】]{0,12}】|[×xX]{2,}|待定|待补充|请填写|TBD", full_text):
        errors.append(f"placeholder left in the document: {m.group(0)!r}")
        break
    if re.search(r"合同|协议", first.text):
        if len(set(re.findall(r"^\s*([甲乙丙丁]方|买方|卖方|出租方|承租方|委托方|受托方)\s*[（(：:]", full_text, re.M))) < 2:
            errors.append("contract: party information for each party (甲方 / 乙方 名称、代码、住所、代表人、联系方式) is missing")
        if not re.search(r"盖章|签字|签章", full_text):
            errors.append("contract: signature block (盖章 / 签字 / 日期) is missing")
        footer_xml = "".join(s.footer._element.xml for s in document.sections)
        if "NUMPAGES" not in footer_xml:
            errors.append("contract: footer needs “第 X 页 共 Y 页” (PAGE and NUMPAGES fields)")
        if re.search(r"(?:人民币|¥|￥)\s*[\d,]{4,}", full_text) and "大写" not in full_text:
            errors.append("contract: amounts need the capitalised form (大写)")
    body_xml = document.element.body.xml
    if "<w:txbxContent" in body_xml:
        errors.append("text boxes used for layout or side notes")
    for section in document.sections:
        header = " ".join(p.text for p in section.header.paragraphs).strip()
        if header and "header" not in ALLOWED:
            errors.append(f"header text the user did not ask for: {header[:40]!r}")
        footer = " ".join(p.text for p in section.footer.paragraphs).strip()
        if footer and not re.fullmatch(r"[\d\s第页共/\-—]*", footer) and "header" not in ALLOWED:
            errors.append(f"footer holds more than a page number: {footer[:40]!r}")


# ---------------------------------------------------------------- PPTX

def check_pptx(path: Path, errors: list[str], warnings: list[str]) -> None:
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        errors.append("python-pptx is required: py -m pip install python-pptx")
        return
    deck = Presentation(str(path))
    slide_area = deck.slide_width * deck.slide_height
    families: set[str] = set()
    for number, slide in enumerate(deck.slides, 1):
        where = f"slide {number}"
        fill = slide.background.fill if slide.background is not None else None
        try:
            if fill is not None and fill.type == 1 and luminance(hex_rgb(str(fill.fore_color.rgb))) < 0.6 and "dark" not in ALLOWED:
                errors.append(f"{where}: dark slide background #{fill.fore_color.rgb}")
        except (AttributeError, TypeError, ValueError):
            pass
        texts = []
        for shape in slide.shapes:
            try:
                if shape.fill.type == 1 and shape.width * shape.height > slide_area * 0.1 and luminance(hex_rgb(str(shape.fill.fore_color.rgb))) < 0.6 and "dark" not in ALLOWED:
                    errors.append(f"{where}: large dark shape #{shape.fill.fore_color.rgb}")
            except (AttributeError, TypeError, ValueError):
                pass
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                text = "".join(run.text for run in paragraph.runs).strip()
                if not text:
                    continue
                texts.append(text)
                check_text(text, where, errors)
                for run in paragraph.runs:
                    if run.font.size is not None and Emu(run.font.size).pt < 16:
                        errors.append(f"{where}: font {Emu(run.font.size).pt:g}pt below 16pt: {run.text[:20]!r}")
                    if run.font.name:
                        families.add(run.font.name.lower())
        if not texts:
            errors.append(f"{where}: slide carries no text conclusion")
            continue
        title = slide.shapes.title.text_frame.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else texts[0]
        if number == 1 and title != texts[0]:
            errors.append(f"{where}: the first line of the cover is not the title: {texts[0][:30]!r}")
        if "accessory" not in ALLOWED and FILLER_SLIDE.match(title):
            errors.append(f"{where}: filler slide {title!r} (pass --allow accessory if the user asked for it)")
        if number == 1 and len(texts) > 1 and "cover" not in ALLOWED:
            errors.append(f"{where}: cover holds more than the title: {texts[1][:30]!r}")
    if len(families) > 1:
        errors.append(f"more than one font family: {sorted(families)}")


# ---------------------------------------------------------------- Markdown

def check_markdown(path: Path, errors: list[str], warnings: list[str]) -> None:
    lines = [line.rstrip() for line in path.read_text(encoding="utf-8").splitlines()]
    content = [(i, line) for i, line in enumerate(lines, 1) if line.strip()]
    if not content:
        errors.append("empty file")
        return
    if not re.match(r"#\s+\S", content[0][1]):
        errors.append(f"line {content[0][0]} must be '# title': {content[0][1][:40]!r}")
    if any(re.match(r"##\s", line) for _, line in content) and len(content) > 1 and not re.match(r"##\s", content[1][1]):
        errors.append(f"line {content[1][0]}: text between the title and the first section: {content[1][1][:40]!r}")
    for number, line in content:
        check_text(line, f"line {number}", errors)
        heading = re.match(r"#{2,4}\s+(.*)", line)
        if heading:
            check_heading(heading.group(1), f"line {number}", errors)
        if line.startswith(">"):
            errors.append(f"line {number}: blockquote used as a side note")
        if re.search(r"!\[[^\]]*\]\([^)]*(?:shields\.io|badge)", line, re.I) or line.strip() in {"[TOC]", "[toc]"}:
            errors.append(f"line {number}: badge or generated table of contents")
    rules = sum(1 for _, line in content if re.fullmatch(r"\s*(?:-{3,}|\*{3,})\s*", line))
    if rules > 2:
        warnings.append(f"{rules} horizontal rules used as decoration")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("files", nargs="+", type=Path)
    parser.add_argument("--allow", action="append", default=[], metavar="KIND",
                        help="interface (prototype pages: interface type scales and the application shell nav, aside, header; add website for a website prototype with 40/48 first-screen titles), or only for items the user explicitly asked for: multi-sheet, chart, accessory, cover, header, "
                             "shell, dark, note, emoji, or a text family (" + ", ".join(FAMILIES) + ")")
    parser.add_argument("--allow-multi-sheet", action="store_true", help="same as --allow multi-sheet")
    args = parser.parse_args(argv)
    ALLOWED.clear()
    ALLOWED.update(kind.strip() for value in args.allow for kind in value.split(","))
    if args.allow_multi_sheet:
        ALLOWED.add("multi-sheet")
    failed = False
    for path in args.files:
        errors: list[str] = []
        warnings: list[str] = []
        suffix = path.suffix.lower()
        if not path.is_file():
            errors.append("file not found")
        elif suffix in {".html", ".htm"}:
            check_html(path, errors, warnings)
        elif suffix in {".xlsx", ".xlsm"}:
            check_xlsx(path, errors, warnings)
        elif suffix == ".docx":
            check_docx(path, errors, warnings)
        elif suffix == ".pptx":
            check_pptx(path, errors, warnings)
        elif suffix in {".md", ".markdown"}:
            check_markdown(path, errors, warnings)
        else:
            errors.append(f"unsupported format {suffix}; review it manually against rules/file-design.html")
        status = "FAIL" if errors else "PASS"
        print(f"{status}: {path} ({len(errors)} error(s), {len(warnings)} warning(s))")
        for item in errors:
            print("  ERROR", item)
        for item in warnings:
            print("  WARN ", item)
        failed |= bool(errors)
    if not failed:
        print("Automated checks passed. Apply the deletion test to every visible element before delivery.")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    raise SystemExit(main())
