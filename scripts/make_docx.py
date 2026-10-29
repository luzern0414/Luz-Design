"""Read the rule card first: py -X utf8 scripts/context.py docx

Build a formatted Word document that follows file-design.html (文字文档 and 文档类型).

    py -X utf8 scripts/make_docx.py spec.json out.docx

spec.json (common keys):
{
  "type": "contract" | "proposal" | "report" | "minutes" | "general",
  "title": "健身房智能门禁采购合同",
  "sections": [                                   # body, in order
    {"heading": "合同标的", "paragraphs": ["..."], "items": ["...", "..."],
     "table": {"header": ["设备", "数量", "单价（元）"], "rows": [["人脸终端", 2, 1680]], "money_cols": [2]},
     "children": [ {"heading": "...", "paragraphs": ["..."]} ]}
  ],
  "attachments": [{"title": "设备清单", "sections": [...]}],  # optional, each starts on a new page
  "style": "paper",                               # optional, only when the user asks for 论文版式 (not for contracts)
  "number_format": "1、"                          # paper only: "1、" (default) or "1."
}
contract only:
  "number": "DLKJ-2026-018",                      # 合同编号, right-aligned under the title
  "parties": [{"role": "甲方", "fields": [["名称", "..."], ["统一社会信用代码", "..."], ["地址", "..."], ["法定代表人", "..."], ["联系人及电话", "..."]]},
              {"role": "乙方", "fields": [...]}],
  "preamble": "甲乙双方根据《中华人民共和国民法典》及相关法律法规，经平等协商，就……达成如下协议：",
  "signature": [{"role": "甲方", "lines": ["（盖章）", "法定代表人或授权代表（签字）："]}, {"role": "乙方", "lines": [...]}]
minutes only:
  "meta": [["时间", "..."], ["地点", "..."], ["主持", "..."], ["参会", "..."], ["记录", "..."]]

Paragraph text may contain {{money:128000}}, which renders as “人民币 128,000.00 元（大写：壹拾贰万捌仟元整）”.
Contracts number clauses 第一条、第二条…, sub-clauses 1.1、1.2; other types number 一、（一）1.（1）.
Footer: contracts “第 X 页 共 Y 页”, other types the page number only.
A table stays on the page of the paragraph that introduces it (short tables whole, long ones with at least two rows).
Table alignment: header centred; money_cols (0-based) show 1,234.00 and align right; a column whose text wraps or runs
past 22 characters aligns left; other columns centre (file-design.html 文字文档 › 表格与图片).

"style": "paper" follows file-design.html 文字文档 › 论文版式: A4 with 2.54 / 3.17 cm margins, 宋体 body and 黑体 headings
with Times New Roman for Latin text, title 18 pt and headings 14 / 12 pt without bold, body 12 pt at 1.5 lines,
headings numbered by Word's own list numbering (1、 or 1. / 1.1 / 1.1.1) instead of typed numbers, black text;
tables keep the ordinary light grid. The built-in check then runs with --allow paper.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

FONT = "微软雅黑"
INK = "262626"
LINE = "D9D9D9"
HEADER_FILL = "F0F5FF"            # 极浅蓝：表头
SIZES = {"title": 22, "h1": 16, "h2": 14, "h3": 12, "body": 12, "table": 10.5, "footer": 10.5}
# 论文版式 (file-design.html 文字文档 › 论文版式). 黑体 and 宋体 have no real bold, so titles and headings stay regular.
PAPER = {"body": "宋体", "head": "黑体", "latin": "Times New Roman", "ink": "000000", "margins_cm": (2.54, 3.17), "edge_cm": 1.5,
         "sizes": {"title": 18, "h1": 14, "h2": 12, "h3": 12, "body": 12, "table": 10.5, "footer": 10.5}}
NUMBER_FORMATS = {"1、": ("%1、", "nothing"), "1.": ("%1.", "space")}   # level-1 text and what follows the number
CN = "零一二三四五六七八九"


def cn_num(n: int) -> str:
    if n <= 10:
        return "十" if n == 10 else CN[n]
    if n < 20:
        return "十" + (CN[n % 10] if n % 10 else "")
    if n < 100:
        return CN[n // 10] + "十" + (CN[n % 10] if n % 10 else "")
    return str(n)


def rmb_upper(amount: float) -> str:
    """人民币金额大写，例如 128000 → 壹拾贰万捌仟元整，100500.5 → 壹拾万零伍佰元伍角整。"""
    digits, units, bigs = "零壹贰叁肆伍陆柒捌玖", ["", "拾", "佰", "仟"], ["", "万", "亿", "万亿"]
    fen = round(amount * 100)
    yuan, jiao, cents = fen // 100, fen // 10 % 10, fen % 10
    text = ""
    if yuan:
        groups, n = [], yuan
        while n:
            groups.append(n % 10000)
            n //= 10000
        gap = False
        for gi in range(len(groups) - 1, -1, -1):
            g = groups[gi]
            if g == 0:
                gap = bool(text)
                continue
            if gap or (text and g < 1000):
                text += "零"
            part, zero = "", False
            for i in range(3, -1, -1):
                d = g // 10 ** i % 10
                if d == 0:
                    zero = bool(part)
                else:
                    if zero:
                        part += "零"
                        zero = False
                    part += digits[d] + units[i]
            text += part + bigs[gi]
            gap = False
        text += "元"
    if not jiao and not cents:
        return (text or "零元") + "整"
    if jiao:
        text += digits[jiao] + "角"
    elif yuan:
        text += "零"
    return text + (digits[cents] + "分" if cents else "整")


def money_text(match: re.Match) -> str:
    value = float(match.group(1))
    return f"人民币 {value:,.2f} 元（大写：{rmb_upper(value)}）"


def build(spec: dict, out: Path) -> Path:
    import docx
    from docx.enum.section import WD_ORIENT  # noqa: F401
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    kind = spec.get("type", "general")
    paper = spec.get("style") == "paper" and kind != "contract"
    sizes = PAPER["sizes"] if paper else SIZES
    ink = PAPER["ink"] if paper else INK
    body_font, head_font, latin = (PAPER["body"], PAPER["head"], PAPER["latin"]) if paper else (FONT, FONT, FONT)
    head_bold = not paper
    doc = docx.Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin = sec.right_margin = Cm(PAPER["margins_cm"][1] if paper else 2.8)
    if paper:
        sec.header_distance = sec.footer_distance = Cm(PAPER["edge_cm"])
    usable = sec.page_width - sec.left_margin - sec.right_margin

    def set_font(run, size, bold=False, east=None):
        run.font.name = latin
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(ink)
        fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            fonts.set(qn(attr), latin)
        fonts.set(qn("w:eastAsia"), east or body_font)

    # Base and heading styles carry the format, so Word's navigation pane and table of contents work.
    normal = doc.styles["Normal"]
    normal.font.name = latin
    normal.font.size = Pt(sizes["body"])
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), body_font)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)
    for level, size in ((1, sizes["h1"]), (2, sizes["h2"]), (3, sizes["h3"])):
        st = doc.styles[f"Heading {level}"]
        st.font.name = latin
        st.font.size = Pt(size)
        st.font.bold = head_bold
        st.font.italic = False
        st.font.color.rgb = RGBColor.from_string(ink)
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), head_font)
        st.paragraph_format.space_before = Pt((24 if paper else 18) if level == 1 else 12)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.line_spacing = 1.0 if paper else 1.5

    def element(tag, **attrs):
        e = OxmlElement(tag)
        for k, v in attrs.items():
            e.set(qn(k), v)
        return e

    # 论文版式: Word's own multilevel list numbers the three heading styles, so no number is typed into a heading.
    paper_num = None
    if paper:
        level_text, level_suffix = NUMBER_FORMATS.get(spec.get("number_format", "1、"), NUMBER_FORMATS["1、"])
        numbering = doc.part.numbering_part.element
        abstract_id = max([int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))], default=-1) + 1
        absn = element("w:abstractNum", **{"w:abstractNumId": str(abstract_id)})
        absn.append(element("w:multiLevelType", **{"w:val": "multilevel"}))
        for ilvl, (text, suffix) in enumerate(((level_text, level_suffix), ("%1.%2", "space"), ("%1.%2.%3", "space"))):
            lvl = element("w:lvl", **{"w:ilvl": str(ilvl)})
            for child in (element("w:start", **{"w:val": "1"}), element("w:numFmt", **{"w:val": "decimal"}),
                          element("w:pStyle", **{"w:val": doc.styles[f"Heading {ilvl + 1}"].style_id}),
                          element("w:suff", **{"w:val": suffix}), element("w:lvlText", **{"w:val": text}),
                          element("w:lvlJc", **{"w:val": "left"})):
                lvl.append(child)
            ppr = element("w:pPr")
            ppr.append(element("w:ind", **{"w:left": "0", "w:firstLine": "0"}))
            lvl.append(ppr)
            rpr = element("w:rPr")
            rpr.append(element("w:rFonts", **{"w:ascii": latin, "w:hAnsi": latin, "w:eastAsia": head_font, "w:cs": latin}))
            rpr.append(element("w:b", **{"w:val": "0"}))
            half_points = str(int(sizes[f"h{ilvl + 1}"] * 2))
            rpr.append(element("w:sz", **{"w:val": half_points}))
            rpr.append(element("w:szCs", **{"w:val": half_points}))
            lvl.append(rpr)
            absn.append(lvl)
        first_num = numbering.find(qn("w:num"))          # every abstractNum precedes the num elements
        if first_num is not None:
            first_num.addprevious(absn)
        else:
            numbering.append(absn)

        def new_num(restart=False):
            num_id = max([int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))], default=0) + 1
            num = element("w:num", **{"w:numId": str(num_id)})
            num.append(element("w:abstractNumId", **{"w:val": str(abstract_id)}))
            if restart:
                override = element("w:lvlOverride", **{"w:ilvl": "0"})
                override.append(element("w:startOverride", **{"w:val": "1"}))
                num.append(override)
            numbering.append(num)
            return num_id

        paper_num = new_num()
        for level in (1, 2, 3):
            numpr = doc.styles[f"Heading {level}"].element.get_or_add_pPr().get_or_add_numPr()
            numpr.get_or_add_ilvl().val = level - 1
            numpr.get_or_add_numId().val = paper_num

    def para(text="", size=None, bold=False, align=None, indent=True, style=None, after=0, east=None):
        p = doc.add_paragraph(style=style)
        if text:
            text = re.sub(r"\{\{money:([\d.]+)\}\}", money_text, str(text))
            set_font(p.add_run(text), size or sizes["body"], bold, east)
        if align:
            p.alignment = align
        if indent and kind != "contract" and not style:
            p.paragraph_format.first_line_indent = Pt((size or sizes["body"]) * 2)
        p.paragraph_format.space_after = Pt(after)
        return p

    def heading(text, level, num_id=None):
        p = doc.add_paragraph(style=f"Heading {level}")
        set_font(p.add_run(text), sizes[f"h{level}"], head_bold, head_font)
        if num_id and num_id != paper_num:            # attachments restart their own numbering
            numpr = p._p.get_or_add_pPr().get_or_add_numPr()
            numpr.get_or_add_ilvl().val = level - 1
            numpr.get_or_add_numId().val = num_id
        return p

    def table(t):
        header, rows = t["header"], t["rows"]
        money = set(t.get("money_cols", []))
        if doc.paragraphs:                             # the introducing paragraph stays with the table
            doc.paragraphs[-1].paragraph_format.keep_with_next = True
        kept = len(rows) if len(rows) <= 8 else 2      # short tables move whole; long ones keep the header and two rows
        tb = doc.add_table(rows=1 + len(rows), cols=len(header))
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), LINE)
            borders.append(e)
        tb._tbl.tblPr.append(borders)
        # Column widths follow content (CJK counts double); cells centre vertically; a 合计 row is bold and shaded.
        def width(v):
            # 微软雅黑西文字母与数字约 0.6 个汉字宽（汉字记 2），按 1.2 估算，避免地址、编号等无空格长串被截断换行。
            return sum(2 if ord(ch) > 0x2E80 else 1.2 for ch in str(v if v is not None else ""))
        weights = [min(max([width(header[c])] + [width(f"{r[c]:,.2f}" if c in money and isinstance(r[c], (int, float)) else r[c]) for r in rows]), 28) + 4
                   for c in range(len(header))]
        tb.autofit = False
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tb._tbl.tblPr.append(layout)
        for c, w in enumerate(weights):
            tb.columns[c].width = int(usable * w / sum(weights))
            for row in tb.rows:
                row.cells[c].width = int(usable * w / sum(weights))
        # 一列只用一种对齐：金额列（money_cols）整列右对齐；其余列任一数据格超过 22 字或在本列宽内会折行，整列左对齐；否则整列居中。
        def _wraps(c):
            capacity = usable * weights[c] / sum(weights) / 12700 / sizes["table"] - 1
            return any(isinstance(v[c], str) and len(v[c]) > min(22, capacity) for v in rows if c < len(v))
        long_cols = {c for c in range(len(header)) if _wraps(c)}
        for r, values in enumerate([header] + rows):
            row = tb.rows[r]
            total_row = r > 0 and str(values[0]).strip() in ("合计", "总计", "小计")
            trpr = row._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            if r == 0:
                trpr.append(OxmlElement("w:tblHeader"))
            for c, v in enumerate(values):
                cell = row.cells[c]
                text = f"{v:,.2f}" if (c in money and isinstance(v, (int, float)) and r) else ("" if v is None else str(v))
                cell.paragraphs[0].text = ""
                set_font(cell.paragraphs[0].add_run(text), sizes["table"], r == 0 or total_row)
                cell.paragraphs[0].paragraph_format.keep_with_next = r < kept
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                # Header centred; amounts right so digits line up; long text left; short values centred (表格对齐).
                long = r and c in long_cols and c not in money
                cell.paragraphs[0].alignment = (WD_ALIGN_PARAGRAPH.CENTER if r == 0 else WD_ALIGN_PARAGRAPH.RIGHT if c in money
                                                else WD_ALIGN_PARAGRAPH.LEFT if long else WD_ALIGN_PARAGRAPH.CENTER)
                cell.paragraphs[0].paragraph_format.line_spacing = 1.25
                if paper:
                    cell.paragraphs[0].paragraph_format.space_before = cell.paragraphs[0].paragraph_format.space_after = Pt(3)
                if r == 0 or total_row:
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), HEADER_FILL if r == 0 else "F5F5F5")
                    cell._tc.get_or_add_tcPr().append(shd)
        para("", indent=False, after=6)

    def body(sections, depth=1, prefix="", num_id=None):
        for i, s in enumerate(sections, 1):
            if kind == "contract":
                label = f"第{cn_num(i)}条 " if depth == 1 else f"{prefix}{i} "
                child_prefix = f"{i}." if depth == 1 else f"{prefix}{i}."
            elif paper:
                label = "" if depth <= 3 else f"（{i}）"   # Word numbers levels 1–3 itself
                child_prefix = ""
            else:
                label = [f"{cn_num(i)}、", f"（{cn_num(i)}）", f"{i}.", f"（{i}）"][min(depth, 4) - 1]
                child_prefix = ""
            if s.get("heading"):
                heading(label + s["heading"], min(depth, 3), num_id)
            for text in s.get("paragraphs", []):
                p = para(text)
                if paper:
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for j, item in enumerate(s.get("items", []), 1):
                mark = f"{child_prefix}{j} " if kind == "contract" else f"{j}. "
                p = para(mark + re.sub(r"\{\{money:([\d.]+)\}\}", money_text, item), indent=False)
                p.paragraph_format.left_indent = Pt(sizes["body"] * 2)
                p.paragraph_format.first_line_indent = Pt(-sizes["body"] * (1.5 if kind != "contract" else 2))
            if s.get("table"):
                table(s["table"])
            if s.get("children"):
                body(s["children"], depth + 1, child_prefix, num_id)

    # Title (and contract number)
    t = para(spec["title"], sizes["title"], head_bold, WD_ALIGN_PARAGRAPH.CENTER, indent=False,
             after=6 if kind == "contract" else (18 if paper else 12), east=head_font)
    t.paragraph_format.keep_with_next = True
    if kind == "contract" and spec.get("number"):
        para(f"合同编号：{spec['number']}", align=WD_ALIGN_PARAGRAPH.RIGHT, indent=False, after=12)
    if kind == "contract":
        for party in spec.get("parties", []):
            p = para(f"{party['role']}：{dict(party['fields']).get('名称', '')}", bold=True, indent=False)
            p.paragraph_format.space_before = Pt(6)
            for k, v in party["fields"]:
                if k != "名称":
                    para(f"{k}：{v}", indent=False)
        if spec.get("preamble"):
            p = para(spec["preamble"], indent=True)
            p.paragraph_format.first_line_indent = Pt(sizes["body"] * 2)
            p.paragraph_format.space_before = Pt(12)
    if kind == "minutes" and spec.get("meta"):
        table({"header": ["项目", "内容"], "rows": spec["meta"]})
    body(spec.get("sections", []))
    if kind == "contract" and spec.get("signature"):
        para("（以下无正文）", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=12)
        sig = spec["signature"]
        st = doc.add_table(rows=1, cols=len(sig))
        for c, party in enumerate(sig):
            cell = st.rows[0].cells[c]
            cell.paragraphs[0].text = ""
            set_font(cell.paragraphs[0].add_run(f"{party['role']}（盖章）："), sizes["body"], True)
            for line in party.get("lines", ["法定代表人或授权代表（签字）：", "日期：　　　　年　　月　　日"]):
                p = cell.add_paragraph()
                p.paragraph_format.space_before = Pt(18)
                set_font(p.add_run(line), sizes["body"])
        st.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for att_i, att in enumerate(spec.get("attachments", []), 1):
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        para(f"附件{cn_num(att_i)}：{att['title']}", sizes["h1"], head_bold, WD_ALIGN_PARAGRAPH.LEFT, indent=False, after=12,
             east=head_font)
        body(att.get("sections", []), num_id=new_num(restart=True) if paper else None)

    # Footer: page numbers only (contracts: 第 X 页 共 Y 页)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def field(paragraph, code):
        for kind_, text in (("begin", None), (None, code), ("separate", None), (None, "1"), ("end", None)):
            r = paragraph.add_run()
            set_font(r, sizes["footer"])
            if kind_:
                fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind_); r._r.append(fc)
            elif text == code:
                it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {code} "; r._r.append(it)
            else:
                r.text = text

    if kind == "contract":
        set_font(fp.add_run("第 "), sizes["footer"]); field(fp, "PAGE"); set_font(fp.add_run(" 页 共 "), sizes["footer"])
        field(fp, "NUMPAGES"); set_font(fp.add_run(" 页"), sizes["footer"])
    else:
        field(fp, "PAGE")
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


def verify_output(out: Path, kind: str, allow: list[str] | None = None) -> int:
    """Run the document checker on what was just written; a builder never hands over a file that fails it."""
    import check_document
    print(f"rule card for this file: py -X utf8 scripts/context.py {kind}  (read it before writing the spec)", file=sys.stderr)
    code = check_document.main([arg for name in allow or [] for arg in ("--allow", name)] + [str(out)])
    if code:
        print(f"FAIL: {out.name} does not meet the {kind} rules above; fix the spec and build again", file=sys.stderr)
    return code


def main(argv: list[str] | None = None) -> int:
    argv = argv if argv is not None else sys.argv[1:]
    if len(argv) != 2:
        print(__doc__)
        return 2
    spec = json.loads(Path(argv[0]).read_text(encoding="utf-8-sig"))
    out = build(spec, Path(argv[1]))
    print("wrote", out)
    paper = spec.get("style") == "paper" and spec.get("type") != "contract"
    return verify_output(out, "docx", ["paper"] if paper else None)


if __name__ == "__main__":
    raise SystemExit(main())
