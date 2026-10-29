"""Rules settled by the 2026-10-29 audit, each with a passing and a failing case: dark fills judged by size and class,
auxiliary-text classes, the website floor, topic words in prose, PDF, amount alignment in files, the publish gates,
deliverable formats, rule-card sizes, the shared-core sync, sibling cards and the tab-icon set."""
from __future__ import annotations

import contextlib
import io
import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import check_document as cd  # noqa: E402
import context  # noqa: E402

from new_session import TAB_ICONS, create_session, pick_icon, tab_icon_href  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
ICON = f'<link rel="icon" href="{tab_icon_href("门店后台")}">'
PAGE = """<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width">
<title>门店后台</title>{icon}<style>body{{font-size:14px;color:#2B2326;background:#fff}} {css}</style></head>
<body><main><h1>{title}</h1><section><h2>概况</h2>{body}</section></main></body></html>"""


def run(body: str, css: str = "", allow: tuple[str, ...] = ("interface",), title: str = "门店后台") -> list[str]:
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "page.html"
        path.write_text(PAGE.format(icon=ICON, css=css, body=body, title=title), encoding="utf-8")
        cd.ALLOWED.clear()
        cd.ALLOWED.update(allow)
        errors: list[str] = []
        try:
            cd.check_html(path, errors, [])
        finally:
            cd.ALLOWED.clear()
    return errors


def having(items: list[str], word: str) -> list[str]:
    return [i for i in items if word in i]


class DarkFillBySizeTest(unittest.TestCase):
    """visual-design.html 检查脚本的认定: a container stays a container whatever its class is called."""

    def test_containers_named_like_small_things_fail(self) -> None:
        for body, css in (('<div class="top-bar">顶栏</div>', ".top-bar{background:#001529;height:56px}"),
                          ('<div class="box-card">卡片</div>', ".box-card{background:#141414;padding:16px}"),
                          ('<div class="status-bar">状态栏</div>', ".status-bar{background:#1f1f1f;height:48px}"),
                          ('<div class="badge">大块</div>', ".badge{background:#1f1f1f;height:40px}")):
            for allow in (("interface",), ()):
                with self.subTest(css=css, allow=allow):
                    self.assertTrue(having(run(body, css, allow), "dark or saturated fill"))

    def test_small_elements_pass(self) -> None:
        for body, css in (('<span class="status-dot"></span><p>营业中</p>', ".status-dot{background:#cf1322;width:8px;height:8px}"),
                          ('<div class="progress-bar"><i></i></div>', ".progress-bar{background:#0958d9;height:6px}"),
                          ('<div class="legend"><i></i>营收</div>', ".legend i{background:#0958d9}"),
                          ('<div class="top-bar">顶栏</div>', ".top-bar{background:#fff}.top-bar::after{background:#1f1f1f}"),
                          ('<span class="new-badge">3</span>', ".new-badge{background:#cf1322;color:#fff}")):
            with self.subTest(css=css):
                self.assertFalse(having(run(body, css), "dark or saturated fill"))

    def test_controls_and_state_rules_use_their_element_and_base_size(self) -> None:
        for body, css in (('<button class="send" type="button">发送</button>', "button.send{width:36px;height:36px;background:#0958d9;color:#fff}"),
                          ('<span class="cbox on"></span>', ".cbox{width:16px;height:16px;background:#fff}.cbox.on{background:#0958d9}"),
                          ('<span class="check" aria-checked="true"><i class="box"></i>同意</span>',
                           '.check .box{flex:0 0 16px;height:16px}.check[aria-checked="true"] .box{background:#0958d9}'),
                          ('<div class="steps"><span class="on">1</span></div>', ".steps span.on{background:#0958d9;color:#fff}")):
            with self.subTest(css=css):
                self.assertFalse(having(run(body, css), "dark or saturated fill"))

    def test_a_large_solid_block_fails_even_with_an_icon_inside(self) -> None:
        css = ".kk{width:44px;height:44px;background:#dd5a12;color:#fff}.kk.on{background:#0958d9}"
        errors = run('<span class="kk on"><svg class="icon"></svg></span>', css)
        self.assertEqual(len(having(errors, "dark or saturated fill")), 2)


class AuxiliaryClassTest(unittest.TestCase):
    def test_words_inside_other_words_do_not_count(self) -> None:
        errors = run('<p class="community">这是一段需要阅读的正文说明</p>', ".community{font-size:12px}")
        self.assertTrue(having(errors, "interface body text 12px"))
        self.assertTrue(having(errors, "12px only for elements"))

    def test_whole_words_and_segments_count(self) -> None:
        for cls in ("time", "price-unit", "order_meta", "tag"):
            with self.subTest(cls=cls):
                self.assertFalse(having(run(f'<p class="{cls}">10-28 09:30</p>', f".{cls}{{font-size:12px}}"), "body text"))


class WebsiteFloorTest(unittest.TestCase):
    def main(self, css: str, *allow: str) -> str:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "site.html"
            path.write_text(PAGE.format(icon=ICON, css=css, body="<p>门店介绍</p>", title="门店官网"), encoding="utf-8")
            out = io.StringIO()
            with contextlib.redirect_stdout(out):
                cd.main([str(path), *[x for a in allow for x in ("--allow", a)]])
            cd.ALLOWED.clear()
        return out.getvalue()

    def test_website_keeps_14px_and_its_own_scale(self) -> None:
        for allow in (("website",), ("interface", "website")):
            with self.subTest(allow=allow):
                self.assertIn("below 14px", self.main(".foot{font-size:12px}", *allow))
                self.assertIn("off the product type scale", self.main(".x{font-size:32px}", *allow))
                self.assertNotIn("ERROR", self.main(".hero{font-size:40px}.foot{font-size:14px}", *allow))

    def test_interface_keeps_12px(self) -> None:
        self.assertNotIn("below", self.main(".time{font-size:12px}", "interface"))


class TopicWordsTest(unittest.TestCase):
    """提示词、官方页面、版本号、SHA-256 are topics in prose and metadata on label lines (visual-design.html 文字纪律 › 边界)."""

    PROSE = ("<p>这周最火的是一套提示词模板，官方页面写着它能把写周报的时间省一半。</p>"
             "<p>新版本号是 <code>3.1</code>，下载包附了 SHA-256 校验值，谁都能核对。</p>")

    def test_titles_and_sentences_pass(self) -> None:
        errors = run(self.PROSE, allow=(), title="一套提示词模板火了")
        self.assertEqual([e for e in errors if "banned" in e], [])

    def test_release_time_in_a_spec_sentence_passes(self) -> None:
        body = "<ul><li>点开版本信息面板写前端版本、后端版本、发布时间与运行环境。</li></ul>"
        self.assertEqual(having(run(body, allow=()), "banned"), [])
        self.assertTrue(having(run("<p>发布时间：2026-10-29</p>", allow=()), "banned"))
        self.assertTrue(having(run("<p>本页更新于今天</p>", allow=()), "banned"))

    def test_label_lines_footers_and_fragments_fail(self) -> None:
        for body in ("<p>版本号：3.1</p>", "<p>提示词：写一份周报</p>", "<p>SHA-256 a1b2c3d4</p>",
                     "<p>本文讨论模板的用法，结论很清楚。</p><footer>官方页面 · 2026</footer>"):
            with self.subTest(body=body):
                self.assertTrue(having(run(body, allow=()), "banned"))

    def test_interface_version_element_is_content(self) -> None:
        body = '<div class="app-version">版本号 1.4.2</div>'
        self.assertFalse(having(run(body), "banned"))
        self.assertTrue(having(run(body, allow=()), "banned"))

    def test_markdown_and_word(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            md = Path(tmp) / "a.md"
            md.write_text("# 一套提示词模板火了\n\n## 发生了什么\n\n官方页面写着它能省一半时间，下载包附了 SHA-256 校验值。\n", encoding="utf-8")
            errors: list[str] = []
            cd.check_markdown(md, errors, [])
            self.assertEqual(having(errors, "banned"), [])
            md.write_text("# 周报模板\n\n## 说明\n\n版本号：3.1\n", encoding="utf-8")
            errors = []
            cd.check_markdown(md, errors, [])
            self.assertTrue(having(errors, "banned"))


class PdfTest(unittest.TestCase):
    def test_pdf_is_skipped_not_failed(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            pdf = Path(tmp) / "a.pdf"
            pdf.write_bytes(b"%PDF-1.4\n")
            out = io.StringIO()
            with contextlib.redirect_stdout(out):
                code = cd.main([str(pdf)])
        self.assertEqual(code, 0)
        self.assertIn("SKIP", out.getvalue())
        self.assertIn("PyMuPDF", out.getvalue())

    def test_unknown_formats_still_fail(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            odd = Path(tmp) / "a.key"
            odd.write_bytes(b"x")
            with contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(cd.main([str(odd)]), 1)


class AmountAlignmentTest(unittest.TestCase):
    """Amounts align right in Excel and Word too; headers centre, long text left, short values centre."""

    def test_xlsx(self) -> None:
        import openpyxl
        from make_xlsx import build
        spec = {"title": "门店月度销售", "columns": [{"name": "区域", "merge": True}, {"name": "门店"},
                                                  {"name": "销售额（元）", "type": "money"}, {"name": "备注", "type": "longtext"}],
                "rows": [["华南", "科技园店", 128000, "新店开业"], ["华南", "海岸城店", 96000, ""]],
                "total": {"label": "合计", "sum": ["销售额（元）"]}}
        with tempfile.TemporaryDirectory() as tmp:
            out = build(spec, Path(tmp) / "a.xlsx")
            ws = openpyxl.load_workbook(out).active
            self.assertEqual(ws.cell(3, 3).alignment.horizontal, "center")      # header
            self.assertEqual(ws.cell(4, 3).alignment.horizontal, "right")       # amount
            self.assertEqual(ws.cell(6, 3).alignment.horizontal, "right")       # total amount
            self.assertEqual(ws.cell(4, 2).alignment.horizontal, "center")      # short value
            self.assertEqual(ws.cell(4, 4).alignment.horizontal, "left")        # long text

    def test_docx(self) -> None:
        import docx
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from make_docx import build
        spec = {"type": "report", "title": "设备采购汇总", "sections": [{"heading": "采购明细", "paragraphs": ["本期采购如下。"],
                "table": {"header": ["设备", "数量", "金额（元）"], "rows": [["人脸终端", 2, 3360], ["合计", "", 3360]], "money_cols": [2]}}]}
        with tempfile.TemporaryDirectory() as tmp:
            out = build(spec, Path(tmp) / "a.docx")
            table = docx.Document(str(out)).tables[0]
            self.assertEqual(table.cell(0, 2).paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.CENTER)
            self.assertEqual(table.cell(1, 2).paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.RIGHT)
            self.assertEqual(table.cell(2, 2).paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.RIGHT)
            self.assertEqual(table.cell(1, 1).paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.CENTER)


class PublishGateTest(unittest.TestCase):
    def test_visual_needs_a_person_approval_of_the_current_board(self) -> None:
        from publish import publish
        from test_session_flow import adopt_rules, approve_artifact, tool_root
        import validate_session
        try:
            with tempfile.TemporaryDirectory() as temp:
                root = tool_root(temp)
                session = create_session(root, "秋季海报", ["visual"], skill_root=ROOT, session_kind="standalone_poster")
                adopt_rules(session, "视觉设计规范 › 三、视觉产出", "安全区与四级文字层级")
                with self.assertRaises(ValueError) as caught:
                    publish(session, "visual")
                self.assertIn("not approved", str(caught.exception))
                approve_artifact(session, "visual", "02-visual/visual-board.html")
                board = session / "raw" / "02-visual" / "visual-board.html"
                board.write_text(board.read_text(encoding="utf-8").replace("</main>", "<p>新增方向说明</p></main>"), encoding="utf-8")
                with self.assertRaises(ValueError) as caught:
                    publish(session, "visual")
                self.assertIn("hash", str(caught.exception))
                approve_artifact(session, "visual", "02-visual/visual-board.html")
                self.assertTrue(publish(session, "visual").is_file())
        finally:
            validate_session.NOTES_FILE = ROOT / "experience" / "notes.json"

    def test_document_page_must_pass_check_document(self) -> None:
        from publish import publish
        from test_session_flow import adopt_rules, tool_root
        import validate_session
        try:
            with tempfile.TemporaryDirectory() as temp:
                root = tool_root(temp)
                session = create_session(root, "门店月度报告", ["document"], skill_root=ROOT)
                adopt_rules(session)
                page = session / "raw" / "03-documents" / "document.html"
                page.write_text(page.read_text(encoding="utf-8").replace("</main>", "<p>数据来源：内部系统</p></main>"), encoding="utf-8")
                with self.assertRaises(ValueError) as caught:
                    publish(session, "document")
                self.assertIn("check_document", str(caught.exception))
        finally:
            validate_session.NOTES_FILE = ROOT / "experience" / "notes.json"


class DeliverableFormatTest(unittest.TestCase):
    def test_slides_pdf_markdown_and_images_may_be_the_final_file(self) -> None:
        from validate_session import validate_session_root
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp) / "007-门店汇报"
            (folder / "raw").mkdir(parents=True)
            for ext in (".pptx", ".pdf", ".md", ".png"):
                final = folder / f"{folder.name}{ext}"
                final.write_bytes(b"x")
                errors: list[str] = []
                validate_session_root(folder, errors)
                self.assertEqual(errors, [], ext)
                final.unlink()
            (folder / "草稿.pptx").write_bytes(b"x")
            errors = []
            validate_session_root(folder, errors)
            self.assertTrue(errors)


class RuleCardSizeTest(unittest.TestCase):
    def test_every_card_carries_8_to_35_kb_of_rule_text(self) -> None:
        for kind in context.card_kinds():
            text, _ = context.make_card(kind, [])
            rules = text.split("## 3 规则（摘自规范页原文）", 1)[1].split("## 4 用什么做", 1)[0]
            size = len(rules.encode("utf-8"))
            with self.subTest(kind=kind):
                self.assertGreaterEqual(size, 8 * 1024, f"{kind}: {size} bytes of rule text is too thin")
                self.assertLessEqual(size, 35_000, f"{kind}: {size} bytes of rule text; narrow its sections")
                self.assertNotIn("已截断", text, f"{kind}: cite whole subsections instead of cutting with max_chars")


class SharedCoreSyncTest(unittest.TestCase):
    def test_default_only_checks_and_write_needs_the_flag(self) -> None:
        import sync_rule_core
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            own = base / "Luz-Design" / "scripts"
            own.mkdir(parents=True)
            shutil.copy(ROOT / "scripts" / "context.py", own / "context.py")
            sibling = base / "Luz-Dev" / "scripts" / "context.py"
            sibling.parent.mkdir(parents=True)
            stale = "# head\n# ---- shared rule-card core (old)\nold = 1\n# ---- end of shared rule-card core\n# tail\n"
            sibling.write_text(stale, encoding="utf-8")
            saved = sync_rule_core.SKILL_ROOT
            sync_rule_core.SKILL_ROOT = base / "Luz-Design"
            try:
                with contextlib.redirect_stdout(io.StringIO()) as out:
                    self.assertEqual(sync_rule_core.main([]), 1)
                self.assertIn("DIFF Luz-Dev", out.getvalue())
                self.assertEqual(sibling.read_text(encoding="utf-8"), stale)
                with contextlib.redirect_stdout(io.StringIO()):
                    self.assertEqual(sync_rule_core.main(["--write"]), 0)
                    self.assertEqual(sync_rule_core.main([]), 0)
                self.assertTrue(sibling.read_text(encoding="utf-8").endswith("# tail\n"))
            finally:
                sync_rule_core.SKILL_ROOT = saved


class SiblingCardTest(unittest.TestCase):
    def test_a_sibling_card_citing_a_missing_heading_fails(self) -> None:
        import validate_skill
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            refs = base / "Luz-Design" / "references"
            refs.mkdir(parents=True)
            shutil.copy(ROOT / "references" / "visual-design.html", refs / "visual-design.html")
            cards_dir = base / "Luz-System" / "references"
            cards_dir.mkdir(parents=True)
            spec = lambda *path: {"page": "visual-design.html", "root": str(refs), "path": ["核心规则", *path]}
            cards = {"cards": {"admin": {"sections": [spec("一、视觉基础", "字体与字号", "小程序紧凑用法")]}}}
            (cards_dir / "cards.json").write_text(json.dumps(cards, ensure_ascii=False), encoding="utf-8")
            saved = validate_skill.ROOT
            validate_skill.ROOT = base / "Luz-Design"
            try:
                self.assertEqual(validate_skill.sibling_card_errors(), [])
                cards["cards"]["admin"]["sections"].append(spec("一、视觉基础", "已改名的小节"))
                (cards_dir / "cards.json").write_text(json.dumps(cards, ensure_ascii=False), encoding="utf-8")
                self.assertTrue(validate_skill.sibling_card_errors())
            finally:
                validate_skill.ROOT = saved


class TabIconSetTest(unittest.TestCase):
    def test_article_pictograms_exist_and_stay_distinct(self) -> None:
        for name in ("bowl", "plane", "book"):
            self.assertIn(name, TAB_ICONS)
        self.assertEqual(len(set(TAB_ICONS.values())), len(TAB_ICONS))
        self.assertEqual(pick_icon("川菜餐饮门店"), "bowl")
        self.assertEqual(pick_icon("新手读书清单"), "book")
        self.assertEqual(pick_icon("会员通知推送"), "plane")
        self.assertEqual(pick_icon("健身房会员"), "dumbbell")


class CutSheetCliTest(unittest.TestCase):
    def test_help_and_missing_arguments(self) -> None:
        import cut_element_sheet
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            with self.assertRaises(SystemExit) as caught:
                cut_element_sheet.main(["--help"])
            self.assertEqual(caught.exception.code, 0)
            with self.assertRaises(SystemExit) as caught:
                cut_element_sheet.main([])
            self.assertEqual(caught.exception.code, 2)


if __name__ == "__main__":
    unittest.main()
