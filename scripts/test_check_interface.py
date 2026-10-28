"""Interface layout checks in check_document.py (--allow interface) and the field-name warning for exported sheets."""
from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import check_document as cd  # noqa: E402

PAGE = """<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width">
<title>原型</title><style>body{{font-size:14px}} {css}</style></head><body><h1>门店后台原型</h1><h2>会员</h2>{body}</body></html>"""


def run(body: str, css: str = "", allow: tuple[str, ...] = ("interface",)) -> tuple[list[str], list[str]]:
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "page.html"
        path.write_text(PAGE.format(css=css, body=body), encoding="utf-8")
        cd.ALLOWED.clear()
        cd.ALLOWED.update(allow)
        errors: list[str] = []
        warnings: list[str] = []
        try:
            cd.check_html(path, errors, warnings)
        finally:
            cd.ALLOWED.clear()
    return errors, warnings


def table(rows: list[list[str]], header: list[str], cls: str = "") -> str:
    head = "".join(f"<th>{h}</th>" for h in header)
    body = "".join("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>" for r in rows)
    return f'<table class="{cls}"><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>'


class BodyTextSizeTest(unittest.TestCase):
    def only(self, items: list[str], word: str) -> list[str]:
        return [e for e in items if word in e]

    def test_small_table_cells_are_errors(self) -> None:
        errors, _ = run(table([["张三", "月卡"]] * 3, ["姓名", "卡种"], "grid"), ".grid td{font-size:13px}")
        self.assertEqual(len(self.only(errors, "body text 13px")), 1)

    def test_inherited_size_counts(self) -> None:
        errors, _ = run('<div class="list"><p>本月新增会员</p><label>门店</label></div>', ".list{font-size:12px}")
        self.assertEqual(len(self.only(errors, "body text 12px")), 2)

    def test_em_and_rem_are_computed(self) -> None:
        errors, _ = run('<div class="box"><p>到店人数</p></div>', ".box{font-size:16px}.box p{font-size:.8em}")
        self.assertTrue(self.only(errors, "body text 12.8px"))
        errors, _ = run('<p class="x">到店人数</p>', ".x{font-size:0.875rem}")
        self.assertFalse(self.only(errors, "body text"))

    def test_auxiliary_classes_may_be_small(self) -> None:
        css = ".time{font-size:12px}.unit-cell{font-size:12px}.meta p{font-size:12px}"
        body = '<table><tr><th>时间</th></tr><tr><td class="time">10-28 09:30</td></tr></table>' \
               '<p class="unit-cell">元</p><div class="meta"><p>更新于今天</p></div>'
        errors, _ = run(body, css)
        self.assertFalse(self.only(errors, "body text"))

    def test_fourteen_px_and_state_rules_pass(self) -> None:
        errors, _ = run("<p>会员总数</p><button>新建</button>", "p{font-size:14px} button:hover{font-size:12px}")
        self.assertFalse(self.only(errors, "body text"))

    def test_more_specific_rule_wins(self) -> None:
        errors, _ = run('<p class="lead">会员总数</p>', "p.lead{font-size:14px} p{font-size:12px}")
        self.assertFalse(self.only(errors, "body text"))
        errors, _ = run('<p class="lead">会员总数</p>', "p{font-size:14px} p.lead{font-size:12px}")
        self.assertTrue(self.only(errors, "body text"))

    def test_mobile_media_query_does_not_decide(self) -> None:
        errors, _ = run("<p>会员总数</p>", "p{font-size:14px} @media (max-width:760px){p{font-size:12px}}")
        self.assertFalse(self.only(errors, "body text"))

    def test_reading_pages_are_not_checked(self) -> None:
        errors, _ = run("<p>会员总数</p>", "p{font-size:16px}", allow=())
        self.assertFalse(self.only(errors, "body text"))


class TitleAndNavTest(unittest.TestCase):
    NAV = '<nav class="side"><a href="#" class="{a}">会员管理</a><a href="#">订单管理</a></nav>'

    def test_title_equal_to_active_item_is_an_error(self) -> None:
        body = self.NAV.format(a="active") + '<div class="page-head"><h3>会员管理</h3></div>'
        errors, _ = run(body)
        self.assertTrue([e for e in errors if "repeats the highlighted navigation item" in e])

    def test_aria_current_counts_as_active(self) -> None:
        body = '<aside><button aria-current="page">会员管理</button></aside><div class="page-title">会员管理</div>'
        errors, _ = run(body)
        self.assertTrue([e for e in errors if "repeats the highlighted navigation item" in e])

    def test_different_title_or_inactive_item_passes(self) -> None:
        for body in (self.NAV.format(a="active") + '<div class="page-head"><h3>会员详情</h3></div>',
                     self.NAV.format(a="") + '<div class="page-head"><h3>会员管理</h3></div>'):
            errors, _ = run(body)
            self.assertFalse([e for e in errors if "repeats the highlighted" in e], body)

    def test_mini_program_tab_bar_is_not_navigation_here(self) -> None:
        body = '<div class="tabbar"><button aria-current="page">我的</button></div><div class="page-title">我的</div>'
        errors, _ = run(body)
        self.assertFalse([e for e in errors if "repeats the highlighted" in e])

    def test_group_label_equal_to_only_item(self) -> None:
        body = '<nav class="menu"><div class="group"><div class="group-title">概览</div><a href="#">概览</a></div>' \
               '<div class="group"><div class="group-title">会员</div><a href="#">会员管理</a><a href="#">会员卡</a></div></nav>'
        errors, _ = run(body)
        found = [e for e in errors if "navigation group" in e]
        self.assertEqual(len(found), 1)
        self.assertIn("概览", found[0])

    def test_label_followed_by_single_item_list(self) -> None:
        body = '<aside><h4>概览</h4><ul><li><a href="#">概览</a></li></ul><h4>营销</h4><ul><li><a href="#">优惠券</a></li></ul></aside>'
        errors, _ = run(body)
        self.assertEqual(len([e for e in errors if "navigation group" in e]), 1)

    def test_group_with_different_single_item_passes(self) -> None:
        body = '<nav class="menu"><div class="group"><div class="group-title">概览</div><a href="#">数据驾驶舱</a></div></nav>'
        errors, _ = run(body)
        self.assertFalse([e for e in errors if "navigation group" in e])


class TableColumnTest(unittest.TestCase):
    def test_empty_column_is_an_error(self) -> None:
        errors, _ = run(table([["张三", "—"], ["李四", "-"], ["王五", ""]], ["姓名", "备注"]))
        self.assertTrue([e for e in errors if "'备注' is empty" in e])

    def test_column_with_one_value_or_controls_passes(self) -> None:
        errors, _ = run(table([["张三", "—"], ["李四", "VIP"], ["王五", "—"]], ["姓名", "备注"]))
        self.assertFalse([e for e in errors if "is empty" in e])
        errors, _ = run(table([["张三", '<button aria-label="编辑"><svg></svg></button>']] * 3, ["姓名", "操作"]))
        self.assertFalse([e for e in errors if "is empty" in e])

    def test_two_rows_are_not_enough(self) -> None:
        errors, _ = run(table([["张三", "—"], ["李四", "—"]], ["姓名", "备注"]))
        self.assertFalse([e for e in errors if "is empty" in e])

    def test_identical_column_is_a_warning(self) -> None:
        rows = [[f"会员{i}", "正常"] for i in range(5)]
        _, warnings = run(table(rows, ["姓名", "状态"]))
        self.assertTrue([w for w in warnings if "'状态' holds '正常' in all 5 rows" in w])

    def test_identical_column_needs_five_rows_and_skips_operations(self) -> None:
        _, warnings = run(table([[f"会员{i}", "正常"] for i in range(4)], ["姓名", "状态"]))
        self.assertFalse([w for w in warnings if "holds" in w])
        _, warnings = run(table([[f"会员{i}", "详情"] for i in range(6)], ["姓名", "操作"]))
        self.assertFalse([w for w in warnings if "holds" in w])

    def test_rowspan_group_cells_are_laid_out(self) -> None:
        body = ('<table><thead><tr><th>模块</th><th>功能</th><th>备注</th></tr></thead><tbody>'
                '<tr><td rowspan="3">会员</td><td>办卡</td><td>—</td></tr><tr><td>续卡</td><td>—</td></tr>'
                '<tr><td>退卡</td><td>—</td></tr></tbody></table>')
        errors, _ = run(body)
        self.assertTrue([e for e in errors if "'备注' is empty" in e])
        self.assertFalse([e for e in errors if "'功能' is empty" in e])

    def test_long_text_padding_warning(self) -> None:
        rows = [["门店", "会员到店后在前台核验人脸并领取手环"]] * 3
        _, warnings = run(table(rows, ["对象", "说明"], "t"), ".t td{padding:8px 12px}")
        self.assertTrue([w for w in warnings if "12px side padding" in w])
        _, warnings = run(table(rows, ["对象", "说明"], "t"), ".t td{padding:12px 16px}")
        self.assertFalse([w for w in warnings if "side padding" in w])
        _, warnings = run(table(rows, ["对象", "说明"], "t"))
        self.assertFalse([w for w in warnings if "side padding" in w], "padding not set: not determinable")

    def test_short_cells_are_not_padding_checked(self) -> None:
        _, warnings = run(table([["张三", "月卡"]] * 3, ["姓名", "卡种"], "t"), ".t td{padding:8px}")
        self.assertFalse([w for w in warnings if "side padding" in w])


class ExportHeaderTest(unittest.TestCase):
    def sheet(self, header: list[str]) -> list[str]:
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl missing")
        book = openpyxl.Workbook()
        ws = book.active
        ws.title = "订单"
        ws["A1"] = "订单导出"
        for k, name in enumerate(header, 1):
            ws.cell(3, k, name)
            ws.cell(4, k, "x")
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "订单.xlsx"
            book.save(path)
            errors: list[str] = []
            warnings: list[str] = []
            cd.check_xlsx(path, errors, warnings)
        return [w for w in warnings if "field names" in w]

    def test_snake_and_camel_case_headers_warn(self) -> None:
        self.assertTrue(self.sheet(["order_id", "客户", "金额（元）"]))
        self.assertTrue(self.sheet(["客户", "createdAt"]))
        self.assertTrue(self.sheet(["客户", "status"]))

    def test_chinese_and_business_abbreviations_pass(self) -> None:
        self.assertFalse(self.sheet(["客户", "金额（元）", "SKU", "GMV（元）"]))


if __name__ == "__main__":
    unittest.main()
