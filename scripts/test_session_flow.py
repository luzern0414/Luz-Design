from __future__ import annotations

import contextlib
import hashlib
import io
import json
import sys
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

from add_mode import add_mode
from new_session import OUTPUT_ROOT, create_session, main as new_session_main
from publish import publish
from validate_session import SCHEMA_VERSION, validate
from validate_skill import validate_rule_page

def tool_root(temp: str) -> Path:
    root = Path(temp) / "sessions"
    root.mkdir()
    return root

def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))

def write_json(path: Path, data: dict) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()

def now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")

def approve_artifact(session: Path, name: str, relative: str) -> None:
    path = session / "raw" / "99-delivery" / "APPROVALS.json"
    data = read_json(path)
    data[name].update({
        "status": "approved", "artifact_path": relative,
        "artifact_sha256": digest(session / "raw" / relative),
        "approved_by": "测试用户", "approved_at": now(), "note": "人工确认",
    })
    write_json(path, data)

def legacy_approvals(**statuses: str) -> dict:
    """APPROVALS.json as written before the Luz-System split (schema 2.0, system stages included)."""
    def record(path: str, status: str = "pending") -> dict:
        return {"status": status, "artifact_path": path, "artifact_sha256": None,
                "approved_at": None, "approved_by": None, "note": None}
    data = {
        "schema_version": "2.0",
        "module_scope": {**record("01-system/01-module-intake.html", statuses.get("module_scope", "pending")),
                         "requested_modules": [], "selected_modules": []},
        "architecture": record("01-system/03-system-architecture.html", statuses.get("architecture", "pending")),
        "system_design": record("01-system/04-system-design.html", statuses.get("system_design", "pending")),
        "visual": record("02-visual/visual-board.html"),
        "prototype": {**record("03-prototype/interactive-prototype.html", statuses.get("prototype", "pending")),
                      "browser_acceptance_path": None, "browser_acceptance_sha256": None},
        "standalone_poster": {"status": "not_requested", "approved_at": None, "approved_by": None, "note": None},
    }
    return data

def legacy_manifest(manifest: dict) -> dict:
    """MANIFEST.json as written before the split: schema 2.0 with the six old stages."""
    stage = lambda path, approval="pending": {"path": path, "sha256": None, "approval": approval}
    return {**manifest, "schema_version": "2.0", "stages": {
        "module_scope": stage("01-system/01-module-intake.html"),
        "research": stage("01-system/02-research.md", "research_pending"),
        "architecture": stage("01-system/03-system-architecture.html"),
        "system_design": stage("01-system/04-system-design.html"),
        "visual": stage("02-visual/visual-board.html"),
        "prototype": stage("03-prototype/interactive-prototype.html"),
    }}

def minimal_rule_page() -> str:
    sections = []
    titles = ["适用范围", "工作流程", "核心规则", "平台矩阵", "案例样本", "输出契约", "质量门禁", "交接规则"]
    ids = ["scope", "workflow", "rules", "platform-matrix", "examples", "output-contract", "quality-gates", "handoff"]
    for index, (section_id, heading) in enumerate(zip(ids, titles), 1):
        icon = f'<svg class="icon-linear" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="1.8" aria-hidden="true" focusable="false"><path d="M{index + 3} 4v16M4 {index + 3}h16"/></svg>'
        sections.append(f'<!-- SECTION:{section_id}:START id={section_id} data-rule-id=visual-{section_id} -->\n<section id="{section_id}" data-rule-id="visual-{section_id}" aria-labelledby="{section_id}-heading"><h2 id="{section_id}-heading" class="section-heading">{icon}<span>{heading}</span></h2><p>按标准执行并验收。</p></section>\n<!-- SECTION:{section_id}:END -->')
    return f'<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><title>视觉设计规范</title><link rel="stylesheet" href="guidelines.css"></head><body><main id="main" aria-label="视觉设计规范"><h1>视觉设计规范</h1>{"".join(sections)}</main></body></html>'

def adopt_rules(session: Path, rule: str = "文件设计规范 › 核心规则 › 所有格式的共同底线", how: str = "第一行即主标题，无说明文字") -> None:
    """Fill BRIEF 采用的规则 and 本次沉淀 the way an agent does after reading the rule card and writing a note."""
    import datetime
    import validate_session
    brief = session / "raw" / "BRIEF.md"
    text = brief.read_text(encoding="utf-8")
    head, tail = text.split("## 本次沉淀", 1)
    head = head.rstrip("\n") + f"\n| {rule} | {how} |\n\n"
    tail = tail.rstrip("\n") + "\n| test-session-note | lesson | 测试会话写入的经验 |\n"
    brief.write_text(head + "## 本次沉淀" + tail, encoding="utf-8")
    notes = session.parent / "notes.json"
    today = datetime.date.today().isoformat()
    notes.write_text(json.dumps({"schema_version": "1.0", "notes": [{"key": "test-session-note", "kind": "lesson", "text": "测试会话写入的经验", "uses": 1, "first_seen": today, "last_confirmed": today}]}, ensure_ascii=False), encoding="utf-8")
    validate_session.NOTES_FILE = notes


class SessionFlowTests(unittest.TestCase):
    def tearDown(self) -> None:
        import validate_session
        validate_session.NOTES_FILE = ROOT / "experience" / "notes.json"  # adopt_rules points it at a temporary store

    def test_h1_first_contract_accepts_only_title_then_section(self) -> None:
        with tempfile.TemporaryDirectory(dir=ROOT) as temp:
            page = Path(temp) / "visual-design.html"
            page.write_text(minimal_rule_page(), encoding="utf-8")
            errors: list[str] = []
            validate_rule_page(page, "visual", 0, "视觉设计规范", errors)
            self.assertEqual(errors, [])

            for bad in (
                minimal_rule_page().replace('<main id="main"', '<header><p>采集日期</p></header><main id="main"'),
                minimal_rule_page().replace('<h1>视觉设计规范</h1>', '<h1>视觉设计规范</h1><p>摘要</p>'),
                minimal_rule_page().replace('id="scope"', 'id="scope"><p>来源与证据状态</p>'),
            ):
                page.write_text(bad, encoding="utf-8")
                errors = []
                validate_rule_page(page, "visual", 0, "视觉设计规范", errors)
                self.assertTrue(errors)

    def test_actual_rules_are_title_first_chinese_and_source_free(self) -> None:
        cases = [("visual-design.html", "visual", "视觉设计规范"), ("file-design.html", "file", "文件设计规范")]
        for filename, prefix, title in cases:
            with self.subTest(filename=filename):
                path = ROOT / "references" / filename
                errors: list[str] = []
                validate_rule_page(path, prefix, 0, title, errors)
                self.assertEqual(errors, [])
                text = path.read_text(encoding="utf-8")
                for forbidden in ("source-card", "<aside", 'id="sources"', "rule-date"):
                    self.assertNotIn(forbidden, text)

    def test_rule_content_covers_ordering_visual_and_spreadsheet_contracts(self) -> None:
        visual = (ROOT / "references/visual-design.html").read_text(encoding="utf-8")
        files = (ROOT / "references/file-design.html").read_text(encoding="utf-8")
        for token in ("扁平线性风格", "取色依据", "字体与字号", "海报", "文字纪律"):
            self.assertIn(token, visual)
        for token in ("只能生成一个工作表", "冻结表头", "数据验证", "打印区域", "不得添加建议"):
            self.assertIn(token, files)

    def test_document_checker_rejects_metadata_and_accepts_clean_html(self) -> None:
        from check_document import check_html
        from new_session import tab_icon_href
        style = "<style>body{background:#fff;color:#202124;font:16px/1.75 \"Microsoft YaHei\",sans-serif}td{font-size:15px}</style>"
        icon = f'<link rel="icon" href="{tab_icon_href("门店月度销售报告")}">'
        with tempfile.TemporaryDirectory() as temp:
            clean = Path(temp) / "clean.html"
            clean.write_text(
                f"<!doctype html><html lang=\"zh-CN\"><head><meta name=\"viewport\" content=\"width=device-width\">{icon}{style}</head>"
                "<body><main><h1>门店月度销售</h1><section><h2>销售概况</h2><table><tr><td>一店</td></tr></table></section></main></body></html>",
                encoding="utf-8",
            )
            errors: list[str] = []
            check_html(clean, errors, [])
            self.assertEqual(errors, [])
            dirty = Path(temp) / "dirty.html"
            dirty.write_text(
                "<!doctype html><html lang=\"zh-CN\"><head><style>aside{background:#20243b}small{font-size:12px}</style></head>"
                "<body><aside>工具名</aside><main><h1>门店月度销售</h1><p>9 个去重来源 · 采集日期 2026-10-26T06:17:05Z</p>"
                "<section><h2>销售概况</h2><p>说明：本表仅供参考</p></section></main></body></html>",
                encoding="utf-8",
            )
            errors = []
            check_html(dirty, errors, [])
            joined = "\n".join(errors)
            for expected in ("no browser-tab icon", "first visible text", "between the h1", "timestamp", "source metadata", "dark or saturated fill", "below 14px", "note or advice"):
                self.assertIn(expected, joined)

    def test_document_checker_catches_problem_families_not_just_listed_words(self) -> None:
        import check_document
        from check_document import check_html, check_xlsx
        import openpyxl
        from openpyxl.styles import Font
        check_document.ALLOWED.clear()
        style = "<style>body{background:#fff;color:#202124;font:16px/1.75 \"Microsoft YaHei\",sans-serif}</style>"
        with tempfile.TemporaryDirectory() as temp:
            page = Path(temp) / "variants.html"
            page.write_text(
                f"<!doctype html><html lang=\"zh-CN\"><head>{style}</head><body><main><h1>门店分析</h1>"
                "<section><h2>概况</h2><p>阅读时长 3 分钟</p><p>希望对您有帮助</p><p data-quoted>原文：以下是最全攻略</p></section>"
                "<section><h2>前言</h2><p>待补充</p></section></main><footer>最后更新 10 月</footer></body></html>",
                encoding="utf-8",
            )
            errors: list[str] = []
            check_html(page, errors, [])
            joined = "\n".join(errors)
            for family in ("production metadata", "chat residue", "accessory section", "placeholder", "application shell"):
                self.assertIn(family, joined)
            self.assertNotIn("最全", joined)
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "门店销售"
            rows = [["门店销售"], [], ["门店", "销售额"], ["一店", 100], ["二店", 200], [], ["口径：含税"]]
            for row in rows:
                sheet.append(row)
            for row in sheet.iter_rows():
                for cell in row:
                    cell.font = Font(name="微软雅黑", sz=16 if cell.row == 1 else 11, b=cell.row in (1, 3))
            sheet["D4"] = "红色表示超标"
            sheet["D4"].font = Font(name="微软雅黑", sz=11)
            sheet.column_dimensions["A"].width = 12
            workbook = Path(temp) / "variants.xlsx"
            book.save(workbook)
            errors = []
            check_xlsx(workbook, errors, [])
            joined = "\n".join(errors)
            self.assertIn("outside the data table", joined)
            self.assertIn("single text cell", joined)

    def test_spreadsheet_builder_passes_and_old_style_sheet_fails(self) -> None:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from check_document import check_xlsx
        from make_xlsx import build
        with tempfile.TemporaryDirectory() as temp:
            good = build({"title": "硬件选型清单", "columns": [{"name": "子系统", "merge": True}, {"name": "设备"},
                                                         {"name": "数量", "type": "number"}],
                          "rows": [["门禁通行", "人脸识别终端", 2], ["门禁通行", "门禁控制器", 1], ["客流统计", "客流相机", 2]]},
                         Path(temp) / "good.xlsx")
            errors: list[str] = []
            check_xlsx(good, errors, [])
            self.assertEqual(errors, [])
            book = Workbook()
            sheet = book.active
            sheet.title = "硬件清单"
            sheet["A1"] = "硬件清单"
            rows = [["序号", "子系统", "设备"], [1, "门禁通行", "人脸识别终端"], [2, "门禁通行", "门禁控制器"], [3, "客流统计", "客流相机"]]
            for r, row in enumerate(rows, 3):
                for c, value in enumerate(row, 1):
                    sheet.cell(r, c, value).font = Font(name="微软雅黑", size=11)
            sheet["A1"].font = Font(name="微软雅黑", size=16, bold=True)
            bad = Path(temp) / "bad.xlsx"
            book.save(bad)
            errors = []
            check_xlsx(bad, errors, [])
            joined = "\n".join(errors)
            self.assertIn("below 12pt", joined)
            self.assertIn("sequence-number column", joined)
            self.assertIn("merge those cells vertically", joined)

    def test_document_checker_catches_small_fonts_foreign_family_and_shared_icons(self) -> None:
        import check_document
        from check_document import check_html
        gift = '<svg class="icon" viewBox="0 0 24 24"><rect x="3" y="8" width="18" height="5" rx="1"/><path d="M5 13v8h14v-8M12 8v13"/></svg>'
        page = ('<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width">'
                '<title>反例</title><style>body{font-family:Arial,"Microsoft YaHei",sans-serif;font-size:14px;color:#1f1f1f;background:#fff}</style>'
                '</head><body><h1>反例</h1><h2>首页</h2><div>' + gift + '<span>新人礼券</span></div><div>' + gift + '<span>每日抽奖</span></div>'
                '<svg viewBox="0 0 9 9"><text font-size="9">坐标</text></svg></body></html>')
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "bad.html"
            path.write_text(page, encoding="utf-8")
            check_document.ALLOWED.clear()
            check_document.ALLOWED.add("interface")
            errors: list[str] = []
            try:
                check_html(path, errors, [])
            finally:
                check_document.ALLOWED.clear()
            joined = "\n".join(errors)
            self.assertIn("'arial'", joined)
            self.assertIn("below 12px", joined)
            self.assertIn("one icon labels different things", joined)

    def test_contract_builder_passes_and_incomplete_contract_fails(self) -> None:
        import docx
        from check_document import check_docx
        from make_docx import build, rmb_upper
        self.assertEqual(rmb_upper(128000), "壹拾贰万捌仟元整")
        self.assertEqual(rmb_upper(100500.5), "壹拾万零伍佰元伍角整")
        party = lambda role, name: {"role": role, "fields": [["名称", name], ["住所", "深圳市南山区"], ["联系人及电话", "张工 0755-1234 5678"]]}
        with tempfile.TemporaryDirectory() as temp:
            good = build({"type": "contract", "title": "设备采购合同", "parties": [party("甲方", "甲公司"), party("乙方", "乙公司")],
                          "preamble": "甲乙双方经平等协商，达成如下协议：",
                          "sections": [{"heading": "价款", "paragraphs": ["合同总价为{{money:4860}}。"]}],
                          "signature": [{"role": "甲方"}, {"role": "乙方"}]}, Path(temp) / "good.docx")
            errors: list[str] = []
            check_docx(good, errors, [])
            self.assertEqual(errors, [])
            bad = Path(temp) / "bad.docx"
            document = docx.Document()
            document.add_heading("设备采购合同", 0)
            document.add_paragraph("甲方：【请填写】")
            document.add_paragraph("合同总价为人民币 4860 元。")
            document.save(bad)
            errors = []
            check_docx(bad, errors, [])
            joined = "\n".join(errors)
            for expected in ("placeholder", "party information", "signature block", "第 X 页 共 Y 页", "大写"):
                self.assertIn(expected, joined)

    def test_paper_style_numbers_headings_in_word_and_needs_allow_paper(self) -> None:
        import docx
        import check_document
        from check_document import check_docx
        from docx.oxml.ns import qn
        from make_docx import build
        spec = {"type": "general", "style": "paper", "title": "项目硬件开发资源申请",
                "sections": [{"heading": "项目信息", "paragraphs": ["门店入口采用 A、B 双门门斗。"],
                              "table": {"header": ["项目", "内容"], "rows": [["协助周期", "一周"]]},
                              "children": [{"heading": "现场条件", "paragraphs": ["二级标题由 Word 编号。"]}]},
                             {"heading": "时间要求", "paragraphs": ["协助周期一周。"]}],
                "attachments": [{"title": "设备清单", "sections": [{"heading": "清单", "paragraphs": ["附件编号重新开始。"]}]}]}
        with tempfile.TemporaryDirectory() as temp:
            out = build(spec, Path(temp) / "paper.docx")
            document = docx.Document(str(out))
            headings = [p for p in document.paragraphs if p.style.name.startswith("Heading")]
            self.assertEqual([p.text for p in headings], ["项目信息", "现场条件", "时间要求", "清单"])   # nothing typed
            self.assertFalse(any(r.bold for p in headings + [document.paragraphs[0]] for r in p.runs))
            numbering = document.part.numbering_part.element
            self.assertIn("%1、", [lvl.get(qn("w:val")) for lvl in numbering.iter(qn("w:lvlText"))])
            self.assertIsNotNone(document.styles["Heading 1"].element.pPr.numPr)
            self.assertIsNotNone(headings[-1]._p.pPr.numPr)                                        # attachment restarts
            fonts = document.styles["Normal"].element.rPr.rFonts
            self.assertEqual((fonts.get(qn("w:ascii")), fonts.get(qn("w:eastAsia"))), ("Times New Roman", "宋体"))
            self.assertAlmostEqual(document.sections[0].left_margin.cm, 3.17, places=2)
            errors: list[str] = []
            check_docx(out, errors, [])
            self.assertTrue(any("more than one font family" in e for e in errors))
            check_document.ALLOWED.clear()
            check_document.ALLOWED.add("paper")
            try:
                errors = []
                check_docx(out, errors, [])
            finally:
                check_document.ALLOWED.clear()
            self.assertEqual(errors, [])

    def test_standalone_poster_session_is_created_validated_and_published(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = tool_root(temp)
            session = create_session(root, "秋季海报", ["visual"], skill_root=ROOT, session_kind="standalone_poster")
            raw = session / "raw"
            approvals = read_json(raw / "99-delivery" / "APPROVALS.json")
            manifest = read_json(raw / "99-delivery" / "MANIFEST.json")
            self.assertEqual(approvals["schema_version"], SCHEMA_VERSION)
            self.assertEqual(set(approvals), {"schema_version", "visual", "document", "standalone_poster"})
            self.assertEqual(approvals["standalone_poster"]["status"], "requested")
            self.assertEqual(manifest["schema_version"], SCHEMA_VERSION)
            self.assertEqual(set(manifest["stages"]), {"visual", "document"})
            self.assertEqual((manifest["session_kind"], manifest["modes"]), ("standalone_poster", ["visual"]))
            for child in ("visual-board.html", "prompts", "references", "generated"):
                self.assertTrue((raw / "02-visual" / child).exists(), child)
            self.assertFalse((raw / "01-system").exists())
            self.assertEqual(validate(session), [])
            approve_artifact(session, "visual", "02-visual/visual-board.html")
            self.assertEqual(validate(session), [])
            self.assertTrue((raw / "00-rules" / "poster.md").is_file())
            with self.assertRaises(ValueError) as caught:
                publish(session, "visual")
            self.assertIn("采用的规则", str(caught.exception))
            adopt_rules(session, "视觉设计规范 › 三、视觉产出", "安全区与四级文字层级")
            final = publish(session, "visual")
            self.assertEqual(final.name, f"{session.name}.html")
            self.assertEqual(validate(session), [])
            # Any change after approval invalidates the hash.
            board = raw / "02-visual" / "visual-board.html"
            board.write_text(board.read_text(encoding="utf-8").replace("方向甲", "方向乙"), encoding="utf-8")
            self.assertTrue(any("hash differs" in error for error in validate(session)))
            # Every task word new_session accepts is also accepted by the validator's BRIEF check.
            key_visual = create_session(root, "春季品牌主视觉", ["visual"], skill_root=ROOT, session_kind="standalone_poster")
            self.assertEqual(validate(key_visual), [])

    def test_visual_mode_is_only_standalone_poster_work(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = tool_root(temp)
            with self.assertRaises(ValueError) as caught:
                create_session(root, "门店首页视觉", ["visual"], skill_root=ROOT)
            self.assertIn("Luz-System", str(caught.exception))
            with self.assertRaises(ValueError):
                create_session(root, "门店首页", ["visual"], skill_root=ROOT, session_kind="standalone_poster")
            document = create_session(root, "月度报告", ["document"], skill_root=ROOT)
            with self.assertRaises(PermissionError):
                add_mode(root, document, "visual")
            self.assertEqual([item.name for item in root.iterdir()], [document.name])

    def test_document_session_is_created_validated_and_published(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = tool_root(temp)
            session = create_session(root, "门店月度报告", ["document"], skill_root=ROOT)
            raw = session / "raw"
            manifest = read_json(raw / "99-delivery" / "MANIFEST.json")
            self.assertEqual((manifest["session_kind"], manifest["modes"]), ("general", ["document"]))
            self.assertTrue((raw / "03-documents" / "document.html").is_file())
            self.assertTrue((raw / "03-documents" / "outputs").is_dir())
            self.assertFalse((raw / "02-visual").exists())
            self.assertEqual(validate(session), [])
            approve_artifact(session, "document", "03-documents/document.html")
            self.assertEqual(validate(session), [])
            self.assertTrue((raw / "00-rules" / "reading-html.md").is_file())
            adopt_rules(session)
            publish(session, "document")
            self.assertEqual(validate(session), [])

            # A spreadsheet or Word file named after the folder is the single final deliverable; anything else is stray.
            sheet_session = create_session(root, "门店销售表", ["document"], skill_root=ROOT, cards=["xlsx"])
            self.assertTrue((sheet_session / "raw" / "00-rules" / "xlsx.md").is_file())
            (sheet_session / f"{sheet_session.name}.xlsx").write_bytes(b"xlsx")
            self.assertTrue(any("采用的规则" in error for error in validate(sheet_session)))
            adopt_rules(sheet_session, "文件设计规范 › 核心规则 › 表格文件", "无序号列，分组名纵向合并")
            self.assertEqual(validate(sheet_session), [])
            (sheet_session / "门店销售表.xlsx").write_bytes(b"xlsx")
            errors = validate(sheet_session)
            self.assertEqual(len(errors), 1)
            self.assertIn("found 门店销售表.xlsx", errors[0])
            (sheet_session / "门店销售表.xlsx").unlink()
            (sheet_session / f"{sheet_session.name}.docx").write_bytes(b"docx")
            self.assertTrue(any("more than one final deliverable" in error for error in validate(sheet_session)))

    def test_system_prototype_and_hardware_are_refused_and_name_luz_system(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = tool_root(temp)
            for mode in ("system", "prototype", "hardware"):
                with self.subTest(mode=mode):
                    with self.assertRaises(ValueError) as caught:
                        create_session(root, "订餐系统", [mode], skill_root=ROOT)
                    self.assertIn("Luz-System", str(caught.exception))
                    stderr = io.StringIO()
                    with contextlib.redirect_stderr(stderr), contextlib.redirect_stdout(io.StringIO()):
                        code = new_session_main(["--name", "订餐系统", "--modes", mode, "--root", str(root)])
                    self.assertEqual(code, 1)
                    self.assertIn("Luz-System", stderr.getvalue())
            self.assertEqual(list(root.iterdir()), [])
            session = create_session(root, "月度报告", ["document"], skill_root=ROOT)
            for mode in ("system", "prototype", "hardware"):
                with self.subTest(add_mode=mode), self.assertRaises(ValueError) as caught:
                    add_mode(root, session, mode)
                self.assertIn("Luz-System", str(caught.exception))
            # System folders or declarations in a session are reported, never validated as Luz Design work.
            (session / "raw" / "01-system").mkdir()
            manifest_path = session / "raw" / "99-delivery" / "MANIFEST.json"
            manifest = read_json(manifest_path)
            write_json(manifest_path, {**manifest, "modes": ["document", "system"]})
            joined = "\n".join(validate(session))
            self.assertIn("01-system/ is system work and belongs to Luz-System", joined)
            self.assertIn("system mode belongs to Luz-System", joined)

    def test_legacy_schema_2_ledgers_are_accepted_and_system_keys_ignored(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = tool_root(temp)
            session = create_session(root, "验收清单", ["document"], skill_root=ROOT)
            delivery = session / "raw" / "99-delivery"
            manifest = read_json(delivery / "MANIFEST.json")
            write_json(delivery / "APPROVALS.json", legacy_approvals())
            write_json(delivery / "MANIFEST.json", legacy_manifest(manifest))
            self.assertEqual(validate(session), [])
            # not_requested system stages, as written for old standalone posters, are ignored too.
            write_json(delivery / "APPROVALS.json", legacy_approvals(module_scope="not_requested", system_design="not_requested"))
            self.assertEqual(validate(session), [])
            # An approved system stage means the session is system work: it is reported, not silently accepted.
            write_json(delivery / "APPROVALS.json", legacy_approvals(system_design="approved"))
            self.assertTrue(any("system_design is approved" in e and "Luz-System" in e for e in validate(session)))
            # Schema 3.0 ledgers must not carry the old system stages.
            legacy = legacy_approvals()
            approvals_3 = {**legacy, "schema_version": SCHEMA_VERSION,
                           "document": {**legacy["visual"], "artifact_path": "03-documents/document.html"}}
            write_json(delivery / "APPROVALS.json", approvals_3)
            write_json(delivery / "MANIFEST.json", manifest)
            stray = [e for e in validate(session) if "must not contain system stages" in e]
            self.assertEqual(len(stray), 1)
            for name in ("module_scope", "architecture", "system_design", "prototype"):
                self.assertIn(name, stray[0])
            # Unknown schema versions are rejected; a legacy system session kind is pointed at Luz-System.
            write_json(delivery / "APPROVALS.json", {**legacy_approvals(), "schema_version": "1.0"})
            write_json(delivery / "MANIFEST.json", {**legacy_manifest(manifest), "session_kind": "system"})
            joined = "\n".join(validate(session))
            self.assertIn("schema_version must be 3.0", joined)
            self.assertIn("system sessions belong to Luz-System", joined)

    @unittest.skipUnless(next(OUTPUT_ROOT.glob("004-*"), None), "legacy 004 session is not on this machine")
    def test_existing_legacy_004_session_has_only_its_stray_spreadsheet_error(self) -> None:
        session = next(OUTPUT_ROOT.glob("004-*"))
        errors = validate(session)
        self.assertEqual(len(errors), 1, errors)
        self.assertIn(".xlsx", errors[0])

if __name__ == "__main__":
    unittest.main(verbosity=2)
